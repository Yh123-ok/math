from datetime import date, datetime
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
from matplotlib import font_manager
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.utils.datetime import from_excel


HERE = Path(__file__).resolve().parent
RESULT = HERE / "result3.xlsx"
ATTACHMENT = HERE / "附件"
OUTPUT = HERE / "figures"
T = 144

# 服务器没有默认中文字体时，明确指定附件环境中已有的字体。
FONT_PATH = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
if Path(FONT_PATH).exists():
    plt.rcParams["font.family"] = font_manager.FontProperties(fname=FONT_PATH).get_name()
plt.rcParams["axes.unicode_minus"] = False

def rows(ws):
    return list(ws.iter_rows(min_row=2, values_only=True))

def day_of(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and value > 30000:
        converted = from_excel(value)
        return converted.date() if hasattr(converted, "date") else converted
    return value

def read_price():
    #读取电价
    wb = openpyxl.load_workbook(ATTACHMENT / "附件1.xlsx", data_only=True, read_only=True)
    price = np.array([float(row[1]) for row in wb.active.iter_rows(min_row=2, values_only=True)])
    wb.close()
    if price.size != T:
        raise ValueError(f"电价长度错误: {price.size}")
    return price

def parse_slot(text):
    start, end = text.split("-")
    #转换下标
    def one(value):
        hour, minute = (int(x) for x in value.split(":")[:2])
        return hour * 6 + minute // 10
    return one(start), one(end)

def emergency_series(emergency_rows, days, price):
    #按照分时电价计算费用
    result = {day: np.zeros(T) for day in days}
    costs = {day: 0.0 for day in days}
    current_day = None
    for row in emergency_rows:
        if row[0] is not None:
            current_day = day_of(row[0])
        if current_day not in result or not row[1] or row[2] is None:
            continue
        start, end = parse_slot(str(row[1]))
        if end > start:
            amount = float(row[2]) / (end - start)
            result[current_day][start:end] += amount
            costs[current_day] += float(5 * amount * price[start:end].sum())
    return result, costs

def month_axis(ax):
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))

def main():
    price = read_price()
    wb = openpyxl.load_workbook(RESULT, data_only=True, read_only=True)
    plan = rows(wb["计划购电量"])
    adjusted = rows(wb["调整购电量"])
    storage = rows(wb["充放电量"])
    emergency = rows(wb["紧急购电量"])
    wb.close()
    if not plan or len(plan) != len(adjusted):
        raise ValueError("result3.xlsx 中计划购电量和调整购电量行数不一致")

    days = [day_of(row[0]) for row in plan]
    plan_values = np.array([row[1:T + 1] for row in plan], dtype=float)
    adjusted_values = np.array([row[1:T + 1] for row in adjusted], dtype=float)
    increase = np.maximum(adjusted_values - plan_values, 0)
    decrease = np.maximum(plan_values - adjusted_values, 0)
    increase_cost = increase * price * 1.5
    refund = decrease * price * 0.5
    emergency_by_day, emergency_cost_by_day = emergency_series(emergency, days, price)
    emergency_totals = np.array([emergency_by_day[day].sum() for day in days])

    # 选一个同时有紧急购电和退款的代表日
    refund_totals = refund.sum(axis=1)
    if emergency_totals.max() > 0 and refund_totals.max() > 0:
        score = (emergency_totals / emergency_totals.max()) * (refund_totals / refund_totals.max())
        chosen_index = int(np.argmax(score))
    else:
        chosen_index = int(np.argmax(emergency_totals))
    chosen_day = days[chosen_index]
    hours = np.arange(T) / 6

    # 图 1：选定日计划曲线 + 全年每日费用。
    fig, axes = plt.subplots(2, 1, figsize=(11, 7), constrained_layout=True)
    axes[0].plot(hours, plan_values[chosen_index], label="计划购电量", linewidth=1.6)
    axes[0].plot(hours, adjusted_values[chosen_index], label="调整购电量", linewidth=1.4)
    axes[0].set_title(f"{chosen_day:%Y年%m月%d日}计划购电与调整购电")
    axes[0].set_xlabel("时刻")
    axes[0].set_ylabel("电量（kWh/10分钟）")
    axes[0].set_xticks(np.arange(0, 25, 4))
    axes[0].grid(alpha=0.25)
    axes[0].legend()
    plan_cost = np.array([row[146] for row in plan], dtype=float)
    normal_cost = np.array([row[146] for row in adjusted], dtype=float)
    axes[1].plot(days, plan_cost, label="计划购电费")
    axes[1].plot(days, normal_cost, label="调整后正常购电费")
    axes[1].set_title("全年每日正常购电费用")
    axes[1].set_xlabel("日期")
    axes[1].set_ylabel("费用（元/天）")
    axes[1].grid(alpha=0.25)
    axes[1].legend()
    month_axis(axes[1])
    fig.savefig(OUTPUT / "购电计划与费用.png", dpi=200, bbox_inches="tight")
    plt.close(fig)

    # 图 2：同图展示计划、紧急购电、调整增购，以及增购费和退款。
    emergency_values = emergency_by_day[chosen_day]
    fig, axes = plt.subplots(2, 1, figsize=(12, 8), sharex=True, constrained_layout=True)
    axes[0].plot(hours, plan_values[chosen_index], label="计划购电量", linewidth=1.7)
    axes[0].plot(hours, adjusted_values[chosen_index], "--", label="调整购电量", linewidth=1.2)
    axes[0].bar(hours, emergency_values, width=1 / 6, alpha=0.55, label="紧急购电量")
    axes[0].bar(hours, increase[chosen_index], width=1 / 6, alpha=0.6, label="调整增购量")
    axes[0].set_title(f"{chosen_day:%Y年%m月%d日}购电构成（紧急购电事件按时段均摊）")
    axes[0].set_ylabel("电量（kWh/10分钟）")
    axes[0].set_xticks(np.arange(0, 25, 2))
    axes[0].grid(alpha=0.25)
    axes[0].legend(ncol=4)
    axes[1].bar(hours - 1 / 12, increase_cost[chosen_index], width=1 / 6,
                alpha=0.65, label="调整增购费用")
    axes[1].bar(hours + 1 / 12, -refund[chosen_index], width=1 / 6,
                alpha=0.65, label="减购退款（负值表示返还）")
    axes[1].axhline(0, color="black", linewidth=0.8)
    axes[1].set_title("调整费用明细")
    axes[1].set_xlabel("时刻")
    axes[1].set_ylabel("金额（元/10分钟）")
    axes[1].grid(alpha=0.25)
    axes[1].legend(ncol=2)
    fig.savefig(OUTPUT / "购电构成与调整费用.png", dpi=200, bbox_inches="tight")
    plt.close(fig)

    # 图 3：选定日四小时储能运行和 SOC。
    storage_by_day = {}
    current_day = None
    for row in storage:
        if row[0] is not None:
            current_day = day_of(row[0])
        if current_day is not None:
            storage_by_day.setdefault(current_day, []).append(row)
    day_storage = storage_by_day.get(chosen_day, [])
    positions = np.arange(len(day_storage))
    fig, axes = plt.subplots(2, 1, figsize=(11, 7), constrained_layout=True)
    axes[0].bar(positions - 0.2, [row[2] or 0 for row in day_storage], 0.4, label="充电量")
    axes[0].bar(positions + 0.2, [row[3] or 0 for row in day_storage], 0.4, label="放电量")
    axes[0].set_xticks(positions, [row[1] for row in day_storage], rotation=25)
    axes[0].set_ylabel("电量（kWh/4小时）")
    axes[0].set_title(f"{chosen_day:%Y年%m月%d日}储能充放电")
    axes[0].legend()
    soc_points = [(row[4], row[5]) for row in day_storage if row[5] is not None]
    axes[1].plot(range(len(soc_points)), [value for _, value in soc_points], marker="o")
    axes[1].set_xticks(range(len(soc_points)), [str(label) for label, _ in soc_points])
    axes[1].set_ylabel("储能量（kWh）")
    axes[1].set_xlabel("时刻")
    axes[1].grid(alpha=0.25)
    fig.savefig(OUTPUT / "储能运行.png", dpi=200, bbox_inches="tight")
    plt.close(fig)

    # 图 4：全年紧急购电量和紧急购电费用。
    fig, axes = plt.subplots(2, 1, figsize=(11, 7), sharex=True, constrained_layout=True)
    axes[0].plot(days, emergency_totals, color="#c44e52")
    axes[0].set_title("全年每日紧急购电量")
    axes[0].set_ylabel("电量（kWh/天）")
    axes[0].grid(alpha=0.25)
    axes[1].plot(days, [emergency_cost_by_day[day] for day in days], color="#dd8452")
    axes[1].set_title("紧急购电费用估算（按事件时段分摊）")
    axes[1].set_xlabel("日期")
    axes[1].set_ylabel("费用（元/天）")
    axes[1].grid(alpha=0.25)
    month_axis(axes[1])
    fig.savefig(OUTPUT / "紧急购电.png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    print(f"图片已写入: {OUTPUT}")
    print(f"购电构成图选取日期: {chosen_day}，紧急购电量 {emergency_totals[chosen_index]:.2f} kWh")

if __name__ == "__main__":
    main()
