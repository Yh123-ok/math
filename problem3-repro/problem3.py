from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
import argparse,csv,json
import numpy as np
from scipy.optimize import linprog
from scipy.sparse import csr_matrix, lil_matrix, vstack
try:
    import openpyxl
except ModuleNotFoundError as exc:
    raise SystemExit("缺少依赖，请先运行: pip install -r requirements.txt") from exc

#1.初始化参数
HERE = Path(__file__).resolve().parent
ROOT = HERE if (HERE / "附件").exists() else HERE.parent
ATTACHMENT = ROOT / "附件"
#（1）时间参数
T = 144                         # 一天 144 个十分钟时段
SLOTS_PER_HOUR,TOTAL_DAYS ,FORECAST_HOURS= 6, 365,24
DT = 1 / SLOTS_PER_HOUR         # 十分钟对应的小时数
ISSUE_HOURS = (0, 6, 12, 18)    # 给出预测的时间点
UPDATE_SLOTS = {36: 1, 72: 2, 108: 3}#在一天的144个时段何时用新光伏预报
SCHEME_UPDATES = {"A": {}, "B": {72: 2}, "C": UPDATE_SLOTS}#对应验证的三种方案
# （2）储能参数
ETA_C, ETA_D = 0.9, 0.9         #储能效率
SOC_MIN, SOC_MAX = 1200.0, 10800.0#限制的电池容量最大最小值
MAX_POWER_KW = 5000.0#限制的最大充放电功率
ENERGY_LIMIT = MAX_POWER_KW * DT
INITIAL_SOC = 6000.0
RESERVE_SOC = SOC_MIN
#（3） 预测参数
BETA ,STARTUP_BETA = 0.5,0.7#b[d,t] = max(历史预测误差的 beta 分位数, 0)
HISTORY_DAYS = 28#预测最多使用前28天数据
HALF_LIFE_DAYS = 7.0 #负荷历史样本的时间衰减半衰期
YESTERDAY_WEIGHT = 0.1#昨日实际负荷与昨日预测负荷之间的残差修正比例
MIN_SAME_WEEKDAY, MIN_SAME_TYPE = 2, 3#采用历史数据时选择的依据，下同
ERROR_WINDOW_DAYS, MIN_ERROR_SAMPLES = 28, 7
# （4）LP 参数和运行阶段
TERMINAL_TARGET, STARTUP_TARGET = 4800.0, 6000.0
RHO_MULTIPLIER = STARTUP_RHO_MULTIPLIER = 1.0
WARMUP_END_INDEX, EVALUATION_START_INDEX = 7, 31
# （5）容差
ADOPTION_TOLERANCE = 1e-7
VALIDATION_TOLERANCE = 1e-5
LP_TOLERANCE = 1e-7
# （6）官方模板中日期、总电量、总费用所在的列（按 1 开始计数）
TOTAL_ENERGY_COLUMN = 2 + T
TOTAL_COST_COLUMN = TOTAL_ENERGY_COLUMN + 1
BLOCK_HOURS = 4
BLOCKS_PER_DAY = 24 // BLOCK_HOURS
SLOTS_PER_BLOCK = BLOCK_HOURS * SLOTS_PER_HOUR
#（7）一天的计划、调整、实时储能和紧急购电结果
@dataclass
class DayResult:
    day: date
    s_start: float
    s_end: float
    q0: np.ndarray               # 00:00 原始计划
    q: np.ndarray                # 最终采用计划
    u: np.ndarray                # 增购量
    v: np.ndarray                # 减购量
    e: np.ndarray                # 紧急购电
    c: np.ndarray                # 充电量
    d: np.ndarray                # 放电量
    w: np.ndarray                # 弃电量
    s: np.ndarray                # SOC，含日初和日末
    adopted_updates: int
#（8）线性规划使用的终端储能参数
@dataclass(frozen=True)
class LPParameters:
    terminal_target: float
    rho: float

# 2.数据读取
def excel_rows(path, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    values = list(ws.iter_rows(values_only=True))
    wb.close()
    return values
#把 Excel 日期和字符串统一为 date
def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("/", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    parts = [int(x) for x in text.split("-")]
    if len(parts) == 3:
        return date(*parts)
    raise ValueError(f"无法解析日期: {value!r}")
#读取附件，换算功率
def read_inputs():
    rows = excel_rows(ATTACHMENT / "附件1.xlsx")
    price = np.array([float(row[1]) for row in rows[1:]], dtype=float)
    cold_load = np.array([float(row[2]) for row in rows[1:]], dtype=float) * DT
    load_rows = excel_rows(ATTACHMENT / "附件2.xlsx", "小区负载")
    pv_rows = excel_rows(ATTACHMENT / "附件2.xlsx", "光伏发电实际功率")
    dates = [as_date(row[0]) for row in load_rows[1:]]
    load_power = np.array([[float(x) for x in row[1:T + 1]] for row in load_rows[1:]])
    pv_power = np.array([[float(x) for x in row[1:T + 1]] for row in pv_rows[1:]])
    forecast = {}
    current_day = None
    for row in excel_rows(ATTACHMENT / "附件3.xlsx")[1:]:
        if row[0] not in (None, ""):
            current_day = as_date(row[0])
        hour = int(str(row[1]).split(":")[0])
        values = np.array([float(x) for x in row[2:2 + FORECAST_HOURS]])
        if current_day is None or values.size != FORECAST_HOURS:
            raise ValueError("附件3预报记录不完整")
        forecast[(current_day, hour)] = values

    load, pv = load_power * DT, pv_power * DT
    if price.size != T or load.shape != (TOTAL_DAYS, T) or pv.shape != (TOTAL_DAYS, T):
        raise ValueError(f"输入维度错误: price={price.shape}, load={load.shape}, pv={pv.shape}")
    if len(forecast) != TOTAL_DAYS * len(ISSUE_HOURS):
        raise ValueError(f"附件3预报记录数错误: {len(forecast)}")
    if np.any(price <= 0):
        raise ValueError("附件1存在非正电价")
    return dates, price, cold_load, pv_power, load, pv, forecast

# 3.负荷预测与裕度
def day_type(day):
    return "saturday" if day.weekday() == 5 else "sunday" if day.weekday() == 6 else "workday"
#（1）负荷预测
def build_load_forecast(dates, load, cold_load):
    result = np.zeros_like(load)
    result[0] = cold_load
    kinds = [day_type(day) for day in dates]
    for d in range(1, len(dates)):
        recent = list(range(max(0, d - HISTORY_DAYS), d))
        exact = [i for i in recent if dates[i].weekday() == dates[d].weekday()]
        same = [i for i in recent if kinds[i] == kinds[d]]
        peers = exact if len(exact) >= MIN_SAME_WEEKDAY else same if len(same) >= MIN_SAME_TYPE else recent
        ages = d - np.array(peers, dtype=float)
        weights = np.power(0.5, ages / HALF_LIFE_DAYS)
        base = np.average(load[peers], axis=0, weights=weights)
        result[d] = np.maximum(base + YESTERDAY_WEIGHT * (load[d - 1] - result[d - 1]), 0)
    return result
#（2）取发布时间之前最近的实测光伏功率
def observed_pv_at_issue(day_index, issue_hour, pv_power):
    if issue_hour == 0:
        return 0.0 if day_index == 0 else float(pv_power[day_index - 1, -1])
    return float(pv_power[day_index, issue_hour * SLOTS_PER_HOUR - 1])
#（3）将逐小时预报线性插值为十分钟电量
def build_pv_forecast(dates, pv_power, forecast_map):
    result = np.zeros((len(dates), len(ISSUE_HOURS), T))
    knot_x = np.arange(25, dtype=float)
    for d, current_day in enumerate(dates):
        for issue_index, issue_hour in enumerate(ISSUE_HOURS):
            knots = np.r_[observed_pv_at_issue(d, issue_hour, pv_power),
                          forecast_map[(current_day, issue_hour)]]
            first = issue_hour * SLOTS_PER_HOUR
            for slot in range(first, T):
                left = (slot - first) / SLOTS_PER_HOUR
                right = (slot + 1 - first) / SLOTS_PER_HOUR
                result[d, issue_index, slot] = max(
                    (np.interp(left, knot_x, knots) + np.interp(right, knot_x, knots)) * DT / 2,
                    0.0,)
    return result
#（4）计算净负荷安全裕度
def build_safety_margins(dates, load, pv, load_hat, pv_hat, beta):
    margins = np.zeros_like(pv_hat)
    errors = np.full_like(pv_hat, np.nan)
    kinds = np.array([day_type(day) for day in dates])
    for d in range(len(dates)):
        history = np.arange(d)
        same = history[kinds[:d] == kinds[d]]
        recent_same = same[same >= d - ERROR_WINDOW_DAYS]
        if recent_same.size >= MIN_ERROR_SAMPLES:
            chosen = recent_same
        elif same.size >= MIN_ERROR_SAMPLES:
            chosen = same
        else:
            chosen = history
        for issue_index, issue_hour in enumerate(ISSUE_HOURS):
            first = issue_hour * SLOTS_PER_HOUR
            if chosen.size:
                margins[d, issue_index, first:] = np.maximum(
                    np.quantile(errors[chosen, issue_index, first:], beta, axis=0), 0)
        # 当天所有发布时间的裕度都算完后，才把当天误差放入历史池。
        actual_net = load[d] - pv[d]
        for issue_index, issue_hour in enumerate(ISSUE_HOURS):
            first = issue_hour * SLOTS_PER_HOUR
            errors[d, issue_index, first:] = (
                actual_net[first:] - (load_hat[d] - pv_hat[d, issue_index])[first:])
    return margins

#4.线性规划
#（1）求解原始购电计划和滚动调整计划
class Scheduler:
    def __init__(self, price):
        self.price = np.asarray(price, dtype=float)
        self.solve_count = 0
    def _solve(self, objective, a_eq, b_eq, a_ub, b_ub, bounds, label):
        result = linprog(
            objective, A_ub=a_ub, b_ub=b_ub, A_eq=a_eq, b_eq=b_eq,
            bounds=bounds, method="highs",options={"primal_feasibility_tolerance": LP_TOLERANCE},
        )
        self.solve_count += 1
        if not result.success:
            raise RuntimeError(f"{label}求解失败: {result.message}")
        return result.x
    def _tie_break(self, x, objective, a_eq, b_eq, a_ub, b_ub, bounds,
                   charge, discharge, waste, emergency, label):
       #主目标相同的解中，优先选不同时充放电的解
        if np.minimum(x[charge], x[discharge]).max(initial=0.0) <= 1e-6:
            return x
        secondary = np.zeros_like(objective)
        secondary[charge] = secondary[discharge] = 1.0
        secondary[waste] = secondary[emergency] = 1e-9
        primary = float(objective @ x)
        a_ub2 = vstack([a_ub, csr_matrix(objective.reshape(1, -1))], format="csr")
        b_ub2 = np.r_[b_ub, primary + max(1e-6, abs(primary) * 1e-9)]
        return self._solve(secondary, a_eq, b_eq, a_ub2, b_ub2, bounds, f"{label}二阶段")
#（2）00:00 原始计划：正常购电、储能和终端 SOC 软约束
    def solve_original(self, demand, initial_soc, params):
        n = len(demand)
        iq, ie = slice(0, n), slice(n, 2 * n)
        ic, id_ = slice(2 * n, 3 * n), slice(3 * n, 4 * n)
        iw, isoc = slice(4 * n, 5 * n), slice(5 * n, 6 * n + 1)
        iz, nvar = 6 * n + 1, 6 * n + 2
        # 变量顺序：q | e | c | d | w | S | z。
        a_eq = lil_matrix((2 * n, nvar))
        for t in range(n):
            a_eq[t, iq.start + t] = 1
            a_eq[t, ie.start + t] = 1
            a_eq[t, ic.start + t] = -1
            a_eq[t, id_.start + t] = 1
            a_eq[t, iw.start + t] = -1
            row = n + t
            a_eq[row, isoc.start + t] = -1
            a_eq[row, isoc.start + t + 1] = 1
            a_eq[row, ic.start + t] = -ETA_C
            a_eq[row, id_.start + t] = 1 / ETA_D
        a_eq = a_eq.tocsr()
        a_ub = lil_matrix((1, nvar))
        a_ub[0, isoc.stop - 1] = a_ub[0, iz] = -1
        a_ub = a_ub.tocsr()
        b_eq = np.r_[demand, np.zeros(n)]
        b_ub = np.array([-params.terminal_target])
        bounds = (
            [(0, None)] * n + [(0, 0)] * n+ [(0, ENERGY_LIMIT)] * n * 2 + [(0, None)] * n
            + [(SOC_MIN, SOC_MAX)] * (n + 1) + [(0, SOC_MAX - SOC_MIN)] )
        bounds[isoc.start] = (initial_soc, initial_soc)
        objective = np.zeros(nvar)
        objective[iq] = self.price[:n]
        objective[ie] = 5 * self.price[:n]
        objective[iz] = params.rho
        x = self._solve(objective, a_eq, b_eq, a_ub, b_ub, bounds, "0:00原计划LP")
        x = self._tie_break(x, objective, a_eq, b_eq, a_ub, b_ub, bounds,
                             ic, id_, iw, ie, "0:00原计划LP")
        return x[iq].copy()
#（3）对尚未执行的时段求“不调整”和“自由调整”两个 LP
    def solve_adjustment(self, q0, current_q, demand, initial_soc, params, fixed):
        n = len(demand)
        start = T - n
        iq, iu = slice(0, n), slice(n, 2 * n)
        iv, ie = slice(2 * n, 3 * n), slice(3 * n, 4 * n)
        ic, id_ = slice(4 * n, 5 * n), slice(5 * n, 6 * n)
        iw, isoc = slice(6 * n, 7 * n), slice(7 * n, 8 * n + 1)
        iz, nvar = 8 * n + 1, 8 * n + 2
        # 变量顺序：q | u | v | e | c | d | w | S | z。
        a_eq = lil_matrix((3 * n, nvar))
        for t in range(n):
            a_eq[t, iq.start + t] = 1
            a_eq[t, iu.start + t] = -1
            a_eq[t, iv.start + t] = 1
            row = n + t
            a_eq[row, iq.start + t] = 1
            a_eq[row, ie.start + t] = 1
            a_eq[row, ic.start + t] = -1
            a_eq[row, id_.start + t] = 1
            a_eq[row, iw.start + t] = -1
            row = 2 * n + t
            a_eq[row, isoc.start + t] = -1
            a_eq[row, isoc.start + t + 1] = 1
            a_eq[row, ic.start + t] = -ETA_C
            a_eq[row, id_.start + t] = 1 / ETA_D
        a_eq = a_eq.tocsr()
        a_ub = lil_matrix((1, nvar))
        a_ub[0, isoc.stop - 1] = a_ub[0, iz] = -1
        a_ub = a_ub.tocsr()
        b_eq = np.r_[q0, demand, np.zeros(n)]
        b_ub = np.array([-params.terminal_target])
        q_bounds = [(float(x), float(x)) for x in current_q] if fixed else [(0, None)] * n
        e_bounds = [(0, None)] * n if fixed else [(0, 0)] * n
        bounds = (
            q_bounds + [(0, None)] * n + [(0, float(x)) for x in q0] + e_bounds
            + [(0, ENERGY_LIMIT)] * n * 2 + [(0, None)] * n
            + [(SOC_MIN, SOC_MAX)] * (n + 1) + [(0, SOC_MAX - SOC_MIN)])
        bounds[isoc.start] = (initial_soc, initial_soc)

        objective = np.zeros(nvar)
        objective[iu] = 1.5 * self.price[start:]
        objective[iv] = -0.5 * self.price[start:]
        objective[ie] = 5 * self.price[start:]
        objective[iz] = params.rho
        label = "不调整候选LP" if fixed else "滚动调整LP"
        x = self._solve(objective, a_eq, b_eq, a_ub, b_ub, bounds, label)
        x = self._tie_break(x, objective, a_eq, b_eq, a_ub, b_ub, bounds,
                             ic, id_, iw, ie, label)
        return x[iq].copy(), float(objective @ x)

#5.逐日仿真
#（1）滚动更新购电计划，并按当天实际负荷、光伏实时执行
def simulate_day(scheduler, d, dates, load, pv, load_hat, pv_hat, margins,
                 initial_soc, params, update_slots=UPDATE_SLOTS):
    q0 = scheduler.solve_original(load_hat[d] - pv_hat[d, 0] + margins[d, 0], initial_soc, params)
    q = q0.copy()
    c = np.zeros(T)
    discharge = np.zeros(T)
    emergency = np.zeros(T)
    waste = np.zeros(T)
    soc = np.zeros(T + 1)
    soc[0] = initial_soc
    adopted = 0
    for slot in range(T):
        if slot in update_slots:
            issue_index = update_slots[slot]
            demand = load_hat[d, slot:] - pv_hat[d, issue_index, slot:] + margins[d, issue_index, slot:]
            _, fixed_cost = scheduler.solve_adjustment(
                q0[slot:], q[slot:], demand, soc[slot], params, True)
            new_q, new_cost = scheduler.solve_adjustment(
                q0[slot:], q[slot:], demand, soc[slot], params, False)
            if fixed_cost - new_cost > ADOPTION_TOLERANCE:
                q[slot:] = new_q
                adopted += 1
        # 缺口先放电，不足部分才紧急购电；富余先充电，多余部分弃电。
        gap = float(load[d, slot] - pv[d, slot] - q[slot])
        if gap > 0:
            discharge[slot] = min(gap, ENERGY_LIMIT, ETA_D * max(soc[slot] - RESERVE_SOC, 0))
            emergency[slot] = max(gap - discharge[slot], 0)
        else:
            c[slot] = min(-gap, ENERGY_LIMIT, max((SOC_MAX - soc[slot]) / ETA_C, 0))
            waste[slot] = max(-gap - c[slot], 0)
        soc[slot + 1] = soc[slot] + ETA_C * c[slot] - discharge[slot] / ETA_D
        if abs(soc[slot + 1] - SOC_MIN) < 1e-8:
            soc[slot + 1] = SOC_MIN
        if abs(soc[slot + 1] - SOC_MAX) < 1e-8:
            soc[slot + 1] = SOC_MAX
    return DayResult(
        dates[d], float(initial_soc), float(soc[-1]), q0, q,
        np.maximum(q - q0, 0), np.maximum(q0 - q, 0), emergency,
        c, discharge, waste, soc, adopted,)
#（2）连续仿真多个日期，上一天的末 SOC 作为下一天初始 SOC
def simulate_range(scheduler, indices, dates, load, pv, load_hat, pv_hat, margins,
                   initial_soc, params, update_slots=UPDATE_SLOTS):
    results, soc = [], float(initial_soc)
    for d in indices:
        result = simulate_day(
            scheduler, d, dates, load, pv, load_hat, pv_hat, margins, soc, params, update_slots)
        results.append(result)
        soc = result.s_end
    return results
#（3）原计划 + 增购费用 - 减购退款 + 紧急购电费用
def settlement(result, price):
    return float(price @ (result.q0 + 1.5 * result.u - 0.5 * result.v + 5 * result.e))
#（4）读取网格搜索冻结的参数；没有文件时使用默认参数
def formal_settings():
    path = HERE / "frozen_parameters.json"
    if not path.exists():
        return BETA, TERMINAL_TARGET, RHO_MULTIPLIER
    values = json.loads(path.read_text(encoding="utf-8"))
    return float(values["safety_quantile"]), float(values["terminal_target_kwh"]), float(values["rho_multiplier"])

# 6.结果校验和 Excel
#（1）检查能量平衡、SOC 连续性和上下界。
def audit_results(results, dates, load, pv):
    day_index = {day: i for i, day in enumerate(dates)}
    max_balance = max_state = max_v_q0 = 0.0
    for previous, current in zip(results, results[1:]):
        max_state = max(max_state, abs(previous.s_end - current.s_start))
    for result in results:
        d = day_index[result.day]
        balance = result.q + result.e + pv[d] + result.d - load[d] - result.c - result.w
        state = result.s[1:] - result.s[:-1] - ETA_C * result.c + result.d / ETA_D
        max_balance = max(max_balance, float(np.max(np.abs(balance))))
        max_state = max(max_state, float(np.max(np.abs(state))))
        max_v_q0 = max(max_v_q0, float(np.max(result.v - result.q0)))
        assert result.e.min() >= -VALIDATION_TOLERANCE
        assert result.s.min() >= SOC_MIN - VALIDATION_TOLERANCE
        assert result.s.max() <= SOC_MAX + VALIDATION_TOLERANCE
        assert result.c.max() <= ENERGY_LIMIT + VALIDATION_TOLERANCE
        assert result.d.max() <= ENERGY_LIMIT + VALIDATION_TOLERANCE
        assert np.isfinite(np.r_[result.q0, result.q, result.e, result.c, result.d, result.s]).all()
    assert max_balance <= VALIDATION_TOLERANCE
    assert max_state <= VALIDATION_TOLERANCE
    assert max_v_q0 <= VALIDATION_TOLERANCE
    return max_balance, max_state, max_v_q0
def copy_style_row(ws, styles, row):
    for column, style in enumerate(styles, 1):
        ws.cell(row, column)._style = copy(style)
def resize_rows(ws, count):
    if ws.max_row > count:
        ws.delete_rows(count + 1, ws.max_row - count)
    elif ws.max_row < count:
        ws.insert_rows(ws.max_row + 1, count - ws.max_row)
def emergency_events(results):
    events = []
    for result in results:
        slot, first = 0, True
        while slot < T:
            if result.e[slot] <= 1e-6:
                slot += 1
                continue
            start = slot
            while slot + 1 < T and result.e[slot + 1] > 1e-6:
                slot += 1
            end = slot + 1
            events.append((
                datetime.combine(result.day, time()) if first else None,
                f"{start // SLOTS_PER_HOUR:02d}:{start % SLOTS_PER_HOUR * 10:02d}-"
                f"{end // SLOTS_PER_HOUR:02d}:{end % SLOTS_PER_HOUR * 10:02d}",
                float(result.e[start:end].sum()),))
            first, slot = False, slot + 1
    return events

def write_workbook(results, price):
    wb = openpyxl.load_workbook(ATTACHMENT / "附件5" / "result3.xlsx")
    output = HERE / "result3.xlsx"
    for sheet, field in (("计划购电量", "q0"), ("调整购电量", "q")):
        ws = wb[sheet]
        resize_rows(ws, len(results) + 1)
        for row_number, result in enumerate(results, 2):
            values = getattr(result, field)
            normal_cost = settlement(result, price) - float(5 * price @ result.e)
            ws.cell(row_number, 1).value = datetime.combine(result.day, time())
            for column, value in enumerate(values, 2):
                ws.cell(row_number, column).value = float(value)
            ws.cell(row_number, TOTAL_ENERGY_COLUMN).value = float(values.sum())
            ws.cell(row_number, TOTAL_COST_COLUMN).value = float(
                price @ result.q0 if field == "q0" else normal_cost)
    # 充放电量按四小时汇总，并保留日初、日末 SOC。
    ws = wb["充放电量"]
    styles = [[copy(ws.cell(row, col)._style) for col in range(1, 7)]
              for row in range(2, 2 + BLOCKS_PER_DAY)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for result in results:
        for block in range(BLOCKS_PER_DAY):
            start, end = block * SLOTS_PER_BLOCK, (block + 1) * SLOTS_PER_BLOCK
            ws.append([
                datetime.combine(result.day, time()) if block == 0 else None,
                f"{block * BLOCK_HOURS}:00-{(block + 1) * BLOCK_HOURS}:00",
                float(result.c[start:end].sum()), float(result.d[start:end].sum()),
                time() if block == 0 else ("24:00" if block == 1 else None),
                result.s_start if block == 0 else result.s_end if block == 1 else None,])
            copy_style_row(ws, styles[block], ws.max_row)
    # 紧急购电按连续时段合并。
    ws = wb["紧急购电量"]
    event_styles = [[copy(ws.cell(row, col)._style) for col in range(1, 4)] for row in (2, 3, 4)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    events = emergency_events(results)
    if events:
        for row_number, event in enumerate(events, 2):
            ws.append(list(event))
            copy_style_row(ws, event_styles[(row_number - 2) % len(event_styles)], row_number)
            if event[0] is not None:
                ws.cell(row_number, 1).number_format = "mm-dd-yy"
    else:
        ws.append([None, None, None])
        copy_style_row(ws, event_styles[0], 2)
    try:
        wb.save(output)
    except PermissionError as exc:
        raise SystemExit(f"无法写入 {output}，请先关闭已打开的 result3.xlsx") from exc
    return output

#7.主流程
def run(smoke_days=None, compare=False):
    dates, price, cold_load, pv_power, load, pv, forecast_map = read_inputs()
    load_hat = build_load_forecast(dates, load, cold_load)
    pv_hat = build_pv_forecast(dates, pv_power, forecast_map)
    beta, terminal_target, rho_multiplier = formal_settings()
    margins = build_safety_margins(dates, load, pv, load_hat, pv_hat, beta)
    startup_margins = build_safety_margins(dates, load, pv, load_hat, pv_hat, STARTUP_BETA)
    scheduler = Scheduler(price)
    average_price = float(price.mean())
    startup_params = LPParameters(STARTUP_TARGET, STARTUP_RHO_MULTIPLIER * average_price)
    formal_params = LPParameters(terminal_target, rho_multiplier * average_price)
    # 预热和校准阶段连续传递 SOC，正式评价从 2 月 1 日开始。
    warmup = simulate_range(
        scheduler, range(WARMUP_END_INDEX), dates, load, pv, load_hat, pv_hat,
        startup_margins, INITIAL_SOC, startup_params,)
    calibration = simulate_range(
        scheduler, range(WARMUP_END_INDEX, EVALUATION_START_INDEX), dates, load, pv,
        load_hat, pv_hat, margins, warmup[-1].s_end, formal_params,)
    end = TOTAL_DAYS if smoke_days is None else min(
        TOTAL_DAYS, EVALUATION_START_INDEX + max(int(smoke_days), 1))
    schemes = SCHEME_UPDATES if compare else {"C": SCHEME_UPDATES["C"]}
    scheme_results = {
        name: simulate_range(
            scheduler, range(EVALUATION_START_INDEX, end), dates, load, pv,
            load_hat, pv_hat, margins, calibration[-1].s_end, formal_params, updates,)
        for name, updates in schemes.items()}
    results = scheme_results["C"]
    if not results:
        raise ValueError("评价期没有结果")
    balance, state, v_q0 = audit_results(results, dates, load, pv)
    output = write_workbook(results, price)
    print(f"已完成 C 方案，结果写入: {output}")
    print(f"评价天数: {len(results)}，原计划费用: {sum(price @ r.q0 for r in results):,.2f} 元")
    print(f"紧急购电费用: {sum(5 * price @ r.e for r in results):,.2f} 元，最终费用: {sum(settlement(r, price) for r in results):,.2f} 元")
    print(f"SOC 起止: {results[0].s_start:.3f} -> {results[-1].s_end:.3f} kWh")
    print(f"校验通过: 能量残差 {balance:.3e}，SOC 残差 {state:.3e}，v-q0 最大值 {v_q0:.3e}")
    print(f"LP 求解次数: {scheduler.solve_count}")
    if compare:
        path = HERE / "scheme_summary.csv"
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(["scheme", "days", "total_cost_yuan", "emergency_energy_kwh",
                             "start_soc_kwh", "end_soc_kwh", "adopted_updates"])
            for name, values in scheme_results.items():
                writer.writerow([
                    name, len(values), sum(settlement(r, price) for r in values),
                    sum(r.e.sum() for r in values), values[0].s_start,
                    values[-1].s_end, sum(r.adopted_updates for r in values),])
        print(f"A/B/C 对比写入: {path}")

def main():
    parser = argparse.ArgumentParser(description="问题三核心求解并写入 result3.xlsx")
    parser.add_argument("--smoke-days", type=int, default=None, help="只运行评价期前 N 天")
    parser.add_argument("--compare", action="store_true", help="同时输出 A/B/C 方案")
    args = parser.parse_args()
    run(args.smoke_days, args.compare)

if __name__ == "__main__":
    main()
