"""读取并核验附件。这里只读取工作簿，不负责写结果文件。"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
import hashlib
import json
import math

import numpy as np
import openpyxl

from config import DELTA_HOURS, PERIODS


@dataclass(frozen=True)
class ProblemData:
    dates: tuple[date, ...]
    source_times: tuple[str, ...]
    interval_labels: tuple[str, ...]
    prices: np.ndarray
    load_kwh: np.ndarray
    pv_kwh: np.ndarray
    audit: dict


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _time_text(value) -> str:
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    return str(value).strip()


def _interval_labels() -> tuple[str, ...]:
    labels = []
    for t in range(PERIODS):
        start = t * 10
        end = start + 10
        labels.append(f"{start//60:02d}:{start%60:02d}-{end//60:02d}:{end%60:02d}")
    return tuple(labels)


def read_and_audit(project_root: Path, output_dir: Path) -> ProblemData:
    attachment1 = project_root / "附件" / "附件1.xlsx"
    attachment2 = project_root / "附件" / "附件2.xlsx"
    template = project_root / "附件" / "附件5" / "result2.xlsx"
    for path in (attachment1, attachment2, template):
        if not path.is_file():
            raise FileNotFoundError(path)

    price_wb = openpyxl.load_workbook(attachment1, read_only=True, data_only=True)
    if price_wb.sheetnames != ["Sheet1"]:
        raise ValueError(f"附件1工作表异常: {price_wb.sheetnames}")
    price_ws = price_wb["Sheet1"]
    if (price_ws.max_row, price_ws.max_column) != (145, 4):
        raise ValueError("附件1应为145行4列")
    expected_headers = ("时间", "电价", "小区负载", "光伏发电预测功率")
    if tuple(c.value for c in price_ws[1]) != expected_headers:
        raise ValueError("附件1表头不匹配")
    price_rows = list(price_ws.iter_rows(values_only=True))
    prices = np.array([row[1] for row in price_rows[1:]], dtype=float)

    wb = openpyxl.load_workbook(attachment2, read_only=True, data_only=True)
    expected_sheets = ["小区负载", "光伏发电实际功率"]
    if wb.sheetnames != expected_sheets:
        raise ValueError(f"附件2工作表异常: {wb.sheetnames}")
    load_ws, pv_ws = wb[expected_sheets[0]], wb[expected_sheets[1]]
    if (load_ws.max_row, load_ws.max_column) != (366, 145):
        raise ValueError("附件2负载表应为366行145列")
    if (pv_ws.max_row, pv_ws.max_column) != (366, 145):
        raise ValueError("附件2光伏表应为366行145列")
    load_rows = list(load_ws.iter_rows(values_only=True))
    pv_rows = list(pv_ws.iter_rows(values_only=True))
    headers_load = tuple(_time_text(x) for x in load_rows[0][1:])
    headers_pv = tuple(_time_text(x) for x in pv_rows[0][1:])
    if headers_load != headers_pv:
        raise ValueError("负载与光伏时间表头不一致")
    expected_times = tuple([f"{m//60:02d}:{m%60:02d}" for m in range(10, 1440, 10)] + ["0:00+1"])
    if headers_load != expected_times:
        raise ValueError("附件2时间轴不是00:10至0:00+1")

    dates = tuple(row[0].date() for row in load_rows[1:])
    pv_dates = tuple(row[0].date() for row in pv_rows[1:])
    expected_dates = tuple(date(2025, 1, 1) + timedelta(days=i) for i in range(365))
    if dates != expected_dates or pv_dates != expected_dates:
        raise ValueError("附件2日期不连续或两表日期不一致")
    load = np.array([row[1:] for row in load_rows[1:]], dtype=float)
    pv = np.array([row[1:] for row in pv_rows[1:]], dtype=float)
    arrays = {"电价": prices, "负载功率": load, "光伏功率": pv}
    for name, array in arrays.items():
        if not np.isfinite(array).all() or (array < 0).any():
            raise ValueError(f"{name}存在缺失、无穷或负值")

    template_wb = openpyxl.load_workbook(template, read_only=False, data_only=False)
    template_shapes = {ws.title: [ws.max_row, ws.max_column] for ws in template_wb.worksheets}
    if template_wb.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise ValueError("result2模板工作表名称异常")
    if template_shapes["计划购电量"] != [335, 147]:
        raise ValueError("result2计划购电表结构异常")
    template_dates = tuple(template_wb["计划购电量"].cell(r, 1).value.date() for r in range(2, 336))
    if template_dates != expected_dates[31:]:
        raise ValueError("result2模板日期不是2025-02-01至2025-12-31")

    audit = {
        "attachment1_sheets": price_wb.sheetnames,
        "attachment1_shape": [price_ws.max_row, price_ws.max_column],
        "attachment1_headers": list(expected_headers),
        "attachment2_sheets": wb.sheetnames,
        "attachment2_shapes": {s: [wb[s].max_row, wb[s].max_column] for s in wb.sheetnames},
        "first_date": dates[0].isoformat(), "last_date": dates[-1].isoformat(),
        "date_count": len(dates), "periods_per_day": PERIODS,
        "first_source_time": headers_load[0], "last_source_time": headers_load[-1],
        "missing_or_nonfinite_count": int(sum((~np.isfinite(x)).sum() for x in arrays.values())),
        "duplicate_date_count": len(dates) - len(set(dates)),
        "ranges": {name: {"min": float(a.min()), "max": float(a.max())} for name, a in arrays.items()},
        "template_sheets": template_wb.sheetnames, "template_shapes": template_shapes,
        "attachment1_sha256": _sha256(attachment1), "attachment2_sha256": _sha256(attachment2),
        "template_sha256": _sha256(template),
        "time_interpretation": "source timestamp is interval end; 00:10 maps to 00:00-00:10",
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "data_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    price_wb.close(); wb.close(); template_wb.close()
    return ProblemData(dates, headers_load, _interval_labels(), prices, load * DELTA_HOURS, pv * DELTA_HOURS, audit)
