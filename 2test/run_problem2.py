"""问题二唯一入口：从任意目录运行，生成全部结果。"""
from __future__ import annotations

import csv
import json
import os
from pathlib import Path
import random
import subprocess
import sys
import time

import numpy as np

BASE = Path(__file__).resolve().parent
PROJECT = BASE.parent
sys.path.insert(0, str(BASE))

import config as cfg
from src.data_io import read_and_audit
from src.forecasting import generate_causal_forecasts
from src.simulation import run_year, replay
from src.validation import validate_results
from src.reporting import build_outputs, make_report
from src.figures import make_figures


def progress(message):
    print(f"[{time.strftime('%H:%M:%S')}] {message}", flush=True)


def ensure_node_modules():
    target = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"
    link = BASE / "node_modules"
    if link.exists():
        return
    if not target.exists():
        raise FileNotFoundError("缺少@oai/artifact-tool；请把可用node_modules链接到2test/node_modules")
    if os.name == "nt":
        def literal(p): return "'" + str(p).replace("'", "''") + "'"
        subprocess.run(["powershell","-NoProfile","-Command",
                        f"New-Item -ItemType Junction -Path {literal(link)} -Target {literal(target)} | Out-Null"],check=True)
    else:
        link.symlink_to(target, target_is_directory=True)


def write_excel():
    ensure_node_modules()
    default_node = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe"
    node = Path(os.environ.get("PROBLEM2_NODE", default_node))
    if not node.exists():
        raise FileNotFoundError("找不到Node.js；请设置PROBLEM2_NODE")
    proc = subprocess.run([str(node), str(BASE/"tools"/"write_result2.mjs"), "write"],
                          capture_output=True, text=True, encoding="utf-8", errors="replace")
    (BASE/"outputs"/"excel_writer.log").write_text(proc.stdout+proc.stderr,encoding="utf-8")
    if proc.returncode:
        raise RuntimeError(proc.stderr[-4000:])
    progress(proc.stdout.strip())


def scheme_comparisons(results, data, forecasts):
    start = data.dates.index(__import__('datetime').date.fromisoformat(cfg.OFFICIAL_START))
    initial_soc = results[start].initial_soc
    rows = []
    for scheme, alpha, kappa in [(0,0.5,1.0),(1,0.5,1.0),(2,0.8,1.0)]:
        replayed, _ = replay(start,len(data.dates),initial_soc,scheme,alpha,kappa,
                             data.load_kwh,data.pv_kwh,data.prices,forecasts)
        rows.append({"scheme":scheme,"alpha":alpha,"kappa":kappa,
                     "total_cost_yuan":sum(r.total_cost for r in replayed),
                     "planned_cost_yuan":sum(r.planned_cost for r in replayed),
                     "emergency_cost_yuan":sum(r.emergency_cost for r in replayed),
                     "emergency_kwh":sum(float(r.actual['emergency'].sum()) for r in replayed),
                     "final_soc_kwh":float(replayed[-1].actual['soc'][-1])})
    official = results[start:]
    rows.append({"scheme":9,"alpha":-1,"kappa":-1,
                 "total_cost_yuan":sum(r.total_cost for r in official),
                 "planned_cost_yuan":sum(r.planned_cost for r in official),
                 "emergency_cost_yuan":sum(r.emergency_cost for r in official),
                 "emergency_kwh":sum(float(r.actual['emergency'].sum()) for r in official),
                 "final_soc_kwh":float(official[-1].actual['soc'][-1])})
    return rows


def validate_excel(data, results):
    import openpyxl
    path = BASE/"result2.xlsx"
    # artifact_tool 生成的工作簿使用流式 XML；普通模式可以可靠取得实际维度。
    wb = openpyxl.load_workbook(path,read_only=False,data_only=True)
    assert wb.sheetnames == ["计划购电量","充放电量","紧急购电量"]
    plan = wb["计划购电量"]
    assert (plan.max_row,plan.max_column)==(335,147)
    official = results[31:]
    max_error=0.0
    for r,result in enumerate(official,2):
        values=np.array([plan.cell(r,c).value for c in range(2,146)],dtype=float)
        max_error=max(max_error,float(np.max(np.abs(values-result.plan['q']))))
        max_error=max(max_error,abs(float(plan.cell(r,146).value)-float(result.plan['q'].sum())))
        max_error=max(max_error,abs(float(plan.cell(r,147).value)-result.total_cost))
    battery=wb["充放电量"]; emergency=wb["紧急购电量"]
    result={"excel_plan_max_abs_error":max_error,"excel_plan_rows":plan.max_row-1,
            "excel_battery_rows":battery.max_row-1,"excel_emergency_rows":emergency.max_row-1,
            "excel_sheet_names_preserved":True}
    assert max_error<=1e-7 and battery.max_row-1==334*6
    wb.close()
    return result


def main():
    random.seed(cfg.SEED); np.random.seed(cfg.SEED)
    for folder in [BASE/"outputs",BASE/"reports",BASE/"figures"]: folder.mkdir(parents=True,exist_ok=True)
    progress("读取并核验附件1、附件2和result2模板")
    data=read_and_audit(PROJECT,BASE/"outputs")
    progress("逐日生成严格因果的负载与光伏预测")
    forecasts=generate_causal_forecasts(data.dates,data.load_kwh,data.pv_kwh)
    progress("开始全年运行与月度滚动参数选择")
    results,selections=run_year(data.dates,data.load_kwh,data.pv_kwh,data.prices,forecasts,progress)
    progress("计算方案0、1、2对照")
    comparisons=scheme_comparisons(results,data,forecasts)
    checks=validate_results(results,data.load_kwh,data.pv_kwh,data.prices,forecasts,BASE/"outputs"/"validation.json")
    daily,emergency,battery=build_outputs(BASE,data,forecasts,results,selections,comparisons,checks)
    make_report(BASE,data,daily,selections,comparisons,checks,emergency)
    progress("填写并渲染result2.xlsx")
    write_excel()
    excel_checks=validate_excel(data,results)
    checks.update(excel_checks)
    progress("生成并校验矢量PDF图")
    checks["vector_pdf_count"] = len(make_figures(BASE))
    (BASE/"outputs"/"validation.json").write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding="utf-8")
    make_report(BASE,data,daily,selections,comparisons,checks,emergency)
    progress("全部计算和验收通过")


if __name__ == "__main__":
    log_path=BASE/"outputs"/"run.log"; log_path.parent.mkdir(parents=True,exist_ok=True)
    class Tee:
        def __init__(self,*streams): self.streams=streams
        def write(self,text):
            for s in self.streams: s.write(text)
        def flush(self):
            for s in self.streams: s.flush()
    with log_path.open("w",encoding="utf-8") as log:
        old_out,old_err=sys.stdout,sys.stderr
        sys.stdout,sys.stderr=Tee(old_out,log),Tee(old_err,log)
        try: main()
        finally: sys.stdout,sys.stderr=old_out,old_err
