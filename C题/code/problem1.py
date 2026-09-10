"""问题一：核验、双时间映射、词典序 LP、模板填写、矢量图与结果验收。

Run from any directory: python <project>/code/problem1.py
Numerical inputs: 附件/附件1.xlsx only. All paths derive from this file.
"""
from __future__ import annotations
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import random
import re
import shutil
import subprocess
import sys
import traceback
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import openpyxl  # read/validate only; authoring uses Artifact Tool + lossless XML patch
from scipy import __version__ as scipy_version
from scipy.optimize import linprog, milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix, csr_matrix, vstack, hstack
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'code' / 'outputs'
QA = ROOT / 'code' / 'qa'
REPORT = ROOT / 'reports'
FIG = ROOT / 'figures'
SOURCE = ROOT / '附件' / '附件1.xlsx'
TEMPLATE = ROOT / '附件' / '附件5' / 'result1.xlsx'
N = 144
DT = 1 / 6
C = 5000 / 6  # exact physical limit, not rounded down
ETA = 0.9
TOL = 1e-6
RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies'
NODE = Path(os.environ.get('PROBLEM1_NODE', RUNTIME / 'node/bin/node.exe'))
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def dump(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_csv(path, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def clock_label(minutes):
    return f'{minutes // 60:02d}:{minutes % 60:02d}'


def interval_label(start, end=None):
    return f'{clock_label(start)}-{clock_label(start + 10 if end is None else end)}'


def minute_of(value):
    if hasattr(value, 'hour'):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    match = re.fullmatch(r'(\d{1,2}):(\d{2})(?:\+(\d+))?', text)
    if not match:
        raise ValueError(f'Cannot parse time: {value!r}')
    h, m, day = match.groups()
    return 60 * int(h) + int(m) + 1440 * int(day or 0)


def parse_interval(label):
    start, end = str(label).split('-')
    a, b = minute_of(start), minute_of(end)
    if b <= a:
        b += 1440
    assert b - a == 10, (label, a, b)
    return a, b


def audit_inputs():
    for p in [ROOT / '建模方案.md', ROOT / 'C题.pdf', SOURCE, TEMPLATE]:
        if not p.is_file():
            raise FileNotFoundError(f'Required input missing: {p}')
    plan = (ROOT / '建模方案.md').read_text(encoding='utf-8-sig')
    for required in ['两个问题共用的基础模型', '问题一：确定性日前最小成本调度', '时间']:
        assert required in plan, required
    pdf_text = '\n'.join(p.extract_text() for p in PdfReader(ROOT / 'C题.pdf').pages)
    (OUT / 'problem1_source_pdf_text.txt').write_text(pdf_text, encoding='utf-8')
    for required in ['12000', '5000', '6000', '1200', '10800', '90%', 'result1.xlsx', '10:00-10:10']:
        assert required in pdf_text, required
    wb = openpyxl.load_workbook(SOURCE, data_only=False)
    assert wb.sheetnames == ['Sheet1'], wb.sheetnames
    sheet = wb['Sheet1']
    values = list(sheet.values)
    headers = ['时间', '电价', '小区负载', '光伏发电预测功率']
    assert list(values[0]) == headers, values[0]
    assert (sheet.max_row, sheet.max_column) == (145, 4)
    raw = values[1:]
    missing = [(r + 2, c + 1) for r, row in enumerate(raw) for c, x in enumerate(row) if x is None]
    nonnumeric = [(r + 2, c + 2, str(x)) for r, row in enumerate(raw) for c, x in enumerate(row[1:])
                  if not isinstance(x, (float, int)) or isinstance(x, bool)]
    times = [minute_of(row[0]) for row in raw]
    if times[-1] == 0:
        times[-1] = 1440
    audit = {'source_sheet_names': wb.sheetnames, 'shape_with_header': [145, 4],
             'numeric_observations': 144, 'headers': headers,
             'first_time': str(raw[0][0]), 'last_time': str(raw[-1][0]),
             'missing_cells': missing, 'nonnumeric_cells': nonnumeric,
             'duplicate_time_count': len(times) - len(set(times)),
             'duplicate_full_row_count': len(raw) - len(set(raw)),
             'input_sha256': sha(SOURCE), 'template_sha256': sha(TEMPLATE),
             'plan_sha256': sha(ROOT / '建模方案.md'), 'pdf_sha256': sha(ROOT / 'C题.pdf'),
             'numpy_version': np.__version__, 'scipy_version': scipy_version,
             'python_version': sys.version.split()[0]}
    assert not missing and not nonnumeric
    assert times == list(range(10, 1441, 10)), times
    numbers = np.asarray([row[1:] for row in raw], dtype=float)
    assert np.isfinite(numbers).all() and (numbers >= 0).all()
    audit['numeric_ranges'] = {headers[j + 1]: {'min': float(numbers[:, j].min()),
                                                        'max': float(numbers[:, j].max()),
                                                        'repeated_numeric_values': 144 - len(set(numbers[:, j]))}
                               for j in range(3)}
    twb = openpyxl.load_workbook(TEMPLATE)
    assert twb.sheetnames == ['计划购电量', '充放电量']
    assert (twb.worksheets[0].max_row, twb.worksheets[0].max_column) == (145, 2)
    assert (twb.worksheets[1].max_row, twb.worksheets[1].max_column) == (7, 5)
    labels = [twb.worksheets[0].cell(i, 1).value for i in range(2, 146)]
    parsed = [parse_interval(s) for s in labels]
    assert parsed == [(t, t + 10) for t in range(10, 1441, 10)]
    audit['template_sheets'] = {s.title: [s.max_row, s.max_column] for s in twb}
    audit['template_first_label'], audit['template_last_label'] = labels[0], labels[-1]
    audit['mapping_A_exact_label_matches'] = len(set(parsed) & {(t - 10, t) for t in times})
    audit['mapping_B_exact_label_matches'] = len(set(parsed) & {(t, t + 10) for t in times})
    assert audit['mapping_A_exact_label_matches'] == 143
    assert audit['mapping_B_exact_label_matches'] == 144
    dump(OUT / 'problem1_data_audit.json', audit)
    print('DATA AUDIT (before optimization)\n' + json.dumps(audit, ensure_ascii=False, indent=2), flush=True)
    return numbers, times, labels, audit


def solve(mapping, numbers, times, labels):
    # Explicit time join. For B, the repeating next-day 00:00 slot belongs to
    # the same civil-time phase as today's 00:00; do not reset SOC at 00:10.
    starts = np.array([t - 10 if mapping == 'A' else t % 1440 for t in times])
    order = np.argsort(starts)
    p, load, pv = numbers[order].T.copy()
    load *= DT
    pv *= DT
    nvar = 4 * N + N + 1
    iq, ic, id_, iw, is_ = [np.arange(k * N, (k + 1) * N) for k in range(4)] + [np.arange(4 * N, nvar)]
    mat = lil_matrix((2 * N, nvar))
    rhs = np.zeros(2 * N)
    for t in range(N):
        mat[t, [iq[t], ic[t], id_[t], iw[t]]] = [1, -1, 1, -1]
        rhs[t] = load[t] - pv[t]
        mat[N + t, [ic[t], id_[t], is_[t], is_[t + 1]]] = [-ETA, 1 / ETA, -1, 1]
    mat = mat.tocsr()
    bounds = [(0, None)] * N + [(0, C)] * (2 * N) + [(0, None)] * N + [(1200, 10800)] * (N + 1)
    bounds[is_[0]] = bounds[is_[-1]] = (6000, 6000)
    cost = np.zeros(nvar)
    cost[iq] = p
    options = {'primal_feasibility_tolerance': 1e-9, 'dual_feasibility_tolerance': 1e-9}
    first = linprog(cost, A_eq=mat, b_eq=rhs, bounds=bounds, method='highs', options=options)
    if not first.success:
        raise RuntimeError(f'{mapping} stage 1 failed: {first.message}')
    eps = max(1e-6, 1e-8 * abs(first.fun))
    throughput = np.zeros(nvar)
    throughput[ic] = throughput[id_] = 1
    second = linprog(throughput, A_eq=mat, b_eq=rhs, A_ub=csr_matrix(cost.reshape(1, -1)),
                     b_ub=[first.fun + eps], bounds=bounds, method='highs', options=options)
    if not second.success:
        raise RuntimeError(f'{mapping} stage 2 failed: {second.message}')
    x = second.x
    method = 'HiGHS LP, lexicographic two stages'
    if np.minimum(x[ic], x[id_]).max() > TOL:
        # Exact physical mutual exclusion, with the same allowed cost tolerance.
        extra = lil_matrix((2 * N, nvar + N))
        for t in range(N):
            extra[t, ic[t]], extra[t, nvar + t] = 1, -C
            extra[N + t, id_[t]], extra[N + t, nvar + t] = 1, C
        a = vstack([hstack([mat, csr_matrix((2 * N, N))]),
                    csr_matrix(np.r_[cost, np.zeros(N)].reshape(1, -1)), extra.tocsr()])
        lo = np.r_[rhs, -np.inf, np.full(2 * N, -np.inf)]
        hi = np.r_[rhs, first.fun + eps, np.zeros(N), np.full(N, C)]
        lower = [b[0] for b in bounds] + [0] * N
        upper = [np.inf if b[1] is None else b[1] for b in bounds] + [1] * N
        integer = milp(np.r_[throughput, np.zeros(N)], integrality=np.r_[np.zeros(nvar), np.ones(N)],
                       bounds=Bounds(lower, upper), constraints=LinearConstraint(a, lo, hi),
                       options={'mip_rel_gap': 1e-10})
        if not integer.success:
            raise RuntimeError(f'{mapping} MILP failed: {integer.message}')
        x = integer.x[:nvar]
        method = 'HiGHS LP cost + MILP throughput with binary mutual exclusion'
    q, c, d, w, s = [x[ix] for ix in [iq, ic, id_, iw, is_]]
    rows = []
    for t, src in enumerate(order):
        rows.append({'mapping': mapping, 'time_label': interval_label(t * 10),
                     'template_time_label': labels[src] if mapping == 'B' else '',
                     'source_excel_row': int(src + 2), 'source_time': clock_label(times[src]),
                     'civil_start_minute': t * 10, 'price_yuan_per_kwh': p[t],
                     'load_kwh': load[t], 'pv_kwh': pv[t], 'purchase_kwh': q[t],
                     'charge_kwh': c[t], 'discharge_kwh': d[t], 'curtailment_kwh': w[t],
                     'soc_start_kwh': s[t], 'soc_end_kwh': s[t + 1]})
    blocks = []
    for h in range(0, 24, 4):
        idx = np.where((starts[order] >= h * 60) & (starts[order] < (h + 4) * 60))[0]
        assert len(idx) == 24
        blocks.append({'mapping': mapping, 'time_label': interval_label(h * 60, (h + 4) * 60),
                       'interval_count': len(idx), 'charge_kwh': float(c[idx].sum()),
                       'discharge_kwh': float(d[idx].sum())})
    baseline = np.maximum(load - pv, 0)
    checks = validate_schedule(rows, blocks)
    # LP dual certificate for the first-stage global optimum.
    dual = float(rhs @ first.eqlin.marginals)
    for i, (lb, ub) in enumerate(bounds):
        dual += lb * first.lower.marginals[i]
        if ub is not None:
            dual += ub * first.upper.marginals[i]
    stationarity = cost - mat.T @ first.eqlin.marginals - first.lower.marginals - first.upper.marginals
    actual_cost, baseline_cost = float(p @ q), float(p @ baseline)
    summary = {'mapping': mapping, 'is_official': mapping == 'B', 'solver': method,
               'stage1_status': int(first.status), 'stage2_status': int(second.status),
               'stage1_message': first.message, 'stage2_message': second.message,
               'optimal_cost_yuan': float(first.fun), 'dispatch_cost_yuan': actual_cost,
               'epsilon_yuan': eps, 'cost_increase_yuan': actual_cost - float(first.fun),
               'cost_bound_violation_yuan': max(0.0, actual_cost - first.fun - eps),
               'lp_dual_gap_yuan': abs(float(first.fun) - dual),
               'lp_stationarity_max': float(np.max(np.abs(stationarity))),
               'purchase_kwh': float(q.sum()), 'charge_kwh': float(c.sum()), 'discharge_kwh': float(d.sum()),
               'curtailment_kwh': float(w.sum()), 'baseline_purchase_kwh': float(baseline.sum()),
               'baseline_cost_yuan': baseline_cost, 'saving_yuan': baseline_cost - actual_cost,
               'saving_fraction': (baseline_cost - actual_cost) / baseline_cost, **checks}
    assert summary['cost_bound_violation_yuan'] <= 1e-6
    assert summary['lp_dual_gap_yuan'] <= 1e-6
    assert summary['lp_stationarity_max'] <= 1e-6
    print(f'SOLVED {mapping}: optimum={first.fun:.10f}, dispatch={actual_cost:.10f}, purchase={q.sum():.10f}', flush=True)
    return rows, blocks, summary


def validate_schedule(rows, blocks):
    def arr(key):
        return np.array([float(row[key]) for row in rows])
    q, c, d, w, s, se, load, pv = [arr(key) for key in ['purchase_kwh', 'charge_kwh', 'discharge_kwh',
                    'curtailment_kwh', 'soc_start_kwh', 'soc_end_kwh', 'load_kwh', 'pv_kwh']]
    values = np.concatenate([q, c, d, w, s, se, load, pv, arr('price_yuan_per_kwh')])
    result = {'energy_balance_max_abs_kwh': float(np.max(np.abs(q + pv + d - load - c - w))),
              'soc_transition_max_abs_kwh': float(np.max(np.abs(se - s - ETA * c + d / ETA))),
              'soc_continuity_max_abs_kwh': float(np.max(np.abs(s[1:] - se[:-1]))),
              'soc_min_kwh': float(min(s.min(), se.min())), 'soc_max_kwh': float(max(s.max(), se.max())),
              'charge_max_kwh': float(c.max()), 'discharge_max_kwh': float(d.max()),
              'soc_0000_kwh': float(s[0]), 'soc_2400_kwh': float(se[-1]),
              'soc_terminal_abs_error_kwh': abs(float(se[-1] - s[0])),
              'simultaneous_charge_discharge_max_kwh': float(np.minimum(c, d).max()),
              'block_charge_sum_abs_error_kwh': abs(sum(float(b['charge_kwh']) for b in blocks) - float(c.sum())),
              'block_discharge_sum_abs_error_kwh': abs(sum(float(b['discharge_kwh']) for b in blocks) - float(d.sum())),
              'minimum_nonnegative_variable_kwh': float(min(q.min(), c.min(), d.min(), w.min())),
              'nan_inf_count': int((~np.isfinite(values)).sum()),
              'period_count': len(rows), 'blocks_each_24': all(int(b['interval_count']) == 24 for b in blocks)}
    for key in ['energy_balance_max_abs_kwh', 'soc_transition_max_abs_kwh', 'soc_continuity_max_abs_kwh',
                'soc_terminal_abs_error_kwh', 'simultaneous_charge_discharge_max_kwh',
                'block_charge_sum_abs_error_kwh', 'block_discharge_sum_abs_error_kwh']:
        assert result[key] <= TOL, (key, result[key])
    assert result['soc_min_kwh'] >= 1200 - TOL and result['soc_max_kwh'] <= 10800 + TOL
    assert c.max() <= C + TOL and d.max() <= C + TOL
    assert result['minimum_nonnegative_variable_kwh'] >= -TOL
    assert result['nan_inf_count'] == 0 and len(rows) == 144 and result['blocks_each_24']
    return result


def node_run(mode):
    if not NODE.exists():
        raise FileNotFoundError('Bundled Node unavailable. Set PROBLEM1_NODE to a Node executable.')
    junction = ROOT / 'code' / 'node_modules'
    if not junction.exists():
        packages = RUNTIME / 'node/node_modules'
        if not packages.exists():
            raise FileNotFoundError('Artifact Tool missing: provide code/node_modules/@oai/artifact-tool')
        if os.name == 'nt':
            # Native PowerShell end-to-end. The path is passed as a quoted literal.
            def ps_literal(p):
                return "'" + str(p).replace("'", "''") + "'"
            cmd = f'New-Item -ItemType Junction -Path {ps_literal(junction)} -Target {ps_literal(packages)} | Out-Null'
            subprocess.run(['powershell', '-NoProfile', '-Command', cmd], check=True)
        else:
            junction.symlink_to(packages, target_is_directory=True)
    proc = subprocess.run([str(NODE), str(ROOT / 'code/fill_template.mjs'), mode],
                          capture_output=True, text=True, encoding='utf-8', errors='replace')
    (OUT / f'artifact_{mode}.log').write_text(proc.stdout + proc.stderr, encoding='utf-8')
    if proc.returncode:
        raise RuntimeError(f'Artifact Tool failed ({mode}): {proc.stderr[-3000:]}')


def fill_template(rows, blocks, labels):
    # Join by the exact original label, not by array row position.
    by_label = {r['template_time_label']: r for r in rows}
    assert len(by_label) == 144 and set(by_label) == set(labels)
    values = {'计划购电量': {f'B{i}': float(by_label[label]['purchase_kwh']) for i, label in enumerate(labels, 2)},
              '充放电量': {}}
    template = openpyxl.load_workbook(TEMPLATE)
    block_lookup = {b['time_label']: b for b in blocks}
    for row in range(2, 8):
        a, b = template['充放电量'].cell(row, 1).value.split('-')
        key = interval_label(minute_of(a), minute_of(b))
        values['充放电量'][f'B{row}'] = float(block_lookup[key]['charge_kwh'])
        values['充放电量'][f'C{row}'] = float(block_lookup[key]['discharge_kwh'])
    values['充放电量']['E2'] = float(rows[0]['soc_start_kwh'])
    values['充放电量']['E3'] = float(rows[-1]['soc_end_kwh'])
    for name, cells in values.items():
        for cell in cells:
            assert template[name][cell].value is None, (name, cell)
    dump(OUT / 'template_values.json', values)
    node_run('fill')
    authored = openpyxl.load_workbook(OUT / 'artifact_filled.xlsx', data_only=True)
    # Restore original package unchanged, transplant only authored numeric payloads.
    # This avoids import/export style normalization: styles, rows, columns, labels,
    # print setup, and every unrelated ZIP member are byte-identical to the template.
    target = ROOT / 'result1.xlsx'
    shutil.copy2(TEMPLATE, target)
    tmp = OUT / 'result1_exact_template.tmp'
    with zipfile.ZipFile(TEMPLATE) as src, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            for idx, name in enumerate(template.sheetnames, 1):
                if info.filename != f'xl/worksheets/sheet{idx}.xml':
                    continue
                xml = content.decode('utf-8')
                for address, expected in values[name].items():
                    val = authored[name][address].value
                    assert isinstance(val, (float, int)) and abs(val - expected) <= 1e-8, (name, address, val, expected)
                    pattern = rf'<c\b([^>/]*\br="{address}"[^>/]*)(?:/>|>.*?</c>)'
                    def replace(match):
                        attrs = re.sub(r'\s+t="[^"]*"', '', match.group(1))
                        return f'<c{attrs} t="n"><v>{format(val, ".17g")}</v></c>'
                    xml, count = re.subn(pattern, replace, xml, flags=re.S)
                    if count == 0:
                        # Excel may omit truly blank/default-style cells from XML.
                        # Materialize only the requested result cell inside its existing row.
                        rownum = re.search(r'\d+', address).group()
                        row_pattern = rf'(<row\b[^>]*\br="{rownum}"[^>]*>)(.*?)(</row>)'
                        new_cell = f'<c r="{address}" t="n"><v>{format(val, ".17g")}</v></c>'
                        def insert_cell(match):
                            body = match.group(2)
                            col = openpyxl.utils.cell.coordinate_from_string(address)[0]
                            colnum = openpyxl.utils.column_index_from_string(col)
                            for existing in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', body):
                                if openpyxl.utils.column_index_from_string(existing.group(1)) > colnum:
                                    body = body[:existing.start()] + new_cell + body[existing.start():]
                                    break
                            else:
                                body += new_cell
                            return match.group(1) + body + match.group(3)
                        xml, count = re.subn(row_pattern, insert_cell, xml, flags=re.S)
                    if count != 1:
                        raise ValueError(f'Cannot fill {name}!{address}: found {count} targets')
                content = xml.encode('utf-8')
            dst.writestr(info, content)
    shutil.copyfile(tmp, target)
    tmp.unlink()
    authored.close()
    (OUT / 'artifact_filled.xlsx').unlink()
    return values


def validate_excel(rows, blocks, values, audit):
    original = openpyxl.load_workbook(TEMPLATE)
    final = openpyxl.load_workbook(ROOT / 'result1.xlsx', data_only=True)
    assert original.sheetnames == final.sheetnames
    max_difference = 0.0
    count = 0
    for name in original.sheetnames:
        a, b = original[name], final[name]
        assert (a.max_row, a.max_column) == (b.max_row, b.max_column)
        for row in a:
            for old in row:
                new = b[old.coordinate]
                assert old.style_id == new.style_id, (name, old.coordinate, 'style')
                if old.coordinate in values[name]:
                    count += 1
                    assert new.data_type == 'n' and math.isfinite(new.value)
                    max_difference = max(max_difference, abs(new.value - values[name][old.coordinate]))
                else:
                    assert old.value == new.value, (name, old.coordinate, 'unrelated value')
    assert count == 158 and max_difference <= 1e-8
    with zipfile.ZipFile(TEMPLATE) as a, zipfile.ZipFile(ROOT / 'result1.xlsx') as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name not in ['xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml']:
                assert a.read(name) == b.read(name), name
            else:
                first, second = ET.fromstring(a.read(name)), ET.fromstring(b.read(name))
                sheetname = original.sheetnames[int(re.search(r'sheet(\d+)', name).group(1)) - 1]
                original_addresses = {cell.attrib['r'] for cell in first.findall('.//m:c', NS)}
                for tree in [first, second]:
                    for row in tree.findall('.//m:row', NS):
                        for cell in list(row):
                            if cell.attrib['r'] in values[sheetname]:
                                if cell.attrib['r'] not in original_addresses:
                                    row.remove(cell)
                                else:
                                    cell.attrib.pop('t', None)
                                    for child in list(cell):
                                        cell.remove(child)
                assert ET.tostring(first) == ET.tostring(second), f'Non-result XML changed: {name}'
    assert sha(TEMPLATE) == audit['template_sha256']
    assert sha(SOURCE) == audit['input_sha256']
    # Independent joins and aggregation from the actual saved CSV and workbook.
    purchase = {r['template_time_label']: float(r['purchase_kwh']) for r in rows}
    errors = [abs(float(final['计划购电量'].cell(i, 2).value) - purchase[final['计划购电量'].cell(i, 1).value])
              for i in range(2, 146)]
    assert max(errors) <= 1e-8
    return {'excel_purchase_max_abs_error_kwh': max(errors), 'excel_numeric_cells_filled': count,
            'excel_all_values_max_abs_error_kwh': max_difference,
            'template_original_unchanged': True, 'source_original_unchanged': True,
            'template_structure_style_labels_preserved': True}


def make_figures():
    # Only read persisted computation results. ReportLab renders vector paths and
    # embedded Chinese TrueType glyphs; no bitmap is placed in any PDF.
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics import renderPDF
    font = Path(os.environ.get('PROBLEM1_CHINESE_FONT', 'C:/Windows/Fonts/simhei.ttf'))
    if not font.exists():
        raise FileNotFoundError('Set PROBLEM1_CHINESE_FONT to a Chinese TrueType font.')
    pdfmetrics.registerFont(TTFont('Chinese', str(font)))
    data = read_csv(OUT / 'problem1_schedule.csv')
    summary = read_csv(OUT / 'problem1_summary.csv')[0]
    def a(key):
        return np.array([float(row[key]) for row in data])
    palette = ['#1D4E89', '#D49A00', '#276D5B', '#B34C46', '#705A9D']
    def panel(can, y, height, series, ymin, ymax, ticks, ylabel, xlabel=False):
        draw = Drawing(730, height + 30)
        lp = LinePlot()
        lp.x, lp.y, lp.width, lp.height = 82, 28, 605, height - 65
        lp.data = [[(float(x), float(yv)) for x, yv in zip(xs, ys)] for _, xs, ys, _ in series]
        lp.xValueAxis.valueMin, lp.xValueAxis.valueMax = 0, 24
        lp.xValueAxis.valueSteps = [0, 4, 8, 12, 16, 20, 24]
        lp.xValueAxis.labelTextFormat = lambda v: f'{int(v):02d}:00'
        lp.yValueAxis.valueMin, lp.yValueAxis.valueMax = ymin, ymax
        lp.yValueAxis.valueSteps = ticks
        lp.yValueAxis.labelTextFormat = '%g'
        for axis in [lp.xValueAxis, lp.yValueAxis]:
            axis.labels.fontName, axis.labels.fontSize = 'Chinese', 9
            axis.strokeColor = colors.HexColor('#718096')
        lp.yValueAxis.visibleGrid = True
        lp.yValueAxis.gridStrokeColor = colors.HexColor('#E5E9EF')
        lp.yValueAxis.gridStrokeWidth = 0.4
        for i, (_, _, _, col) in enumerate(series):
            lp.lines[i].strokeColor = colors.HexColor(col)
            lp.lines[i].strokeWidth = 1.25
        draw.add(lp)
        renderPDF.draw(draw, can, 0, y)
        can.saveState()
        can.translate(22, y + 28 + (height - 65) / 2)
        can.rotate(90)
        can.setFont('Chinese', 11)
        can.drawCentredString(0, 0, ylabel)
        can.restoreState()
        can.setFont('Chinese', 10)
        xx = 90
        for name, _, _, col in series:
            can.setStrokeColor(colors.HexColor(col))
            can.setLineWidth(1.6)
            can.line(xx, y + height - 16, xx + 18, y + height - 16)
            can.setFillColor(colors.HexColor('#263445'))
            can.drawString(xx + 23, y + height - 19, name)
            xx += pdfmetrics.stringWidth(name, 'Chinese', 10) + 44
        if xlabel:
            can.drawCentredString(385, y - 5, '日内时刻')
    xs = np.repeat(np.arange(145) / 6, 2)[1:-1]
    def stepped(key):
        return np.repeat(a(key), 2)
    soc = np.r_[a('soc_start_kwh'), a('soc_end_kwh')[-1]]
    sx = np.arange(145) / 6
    can = canvas.Canvas(str(FIG / 'problem1_dispatch.pdf'), pagesize=(730, 760), invariant=1)
    price_max = math.ceil(float(a('price_yuan_per_kwh').max()) / .3) * .3
    panel(can, 510, 220, [('电价', xs, stepped('price_yuan_per_kwh'), palette[0])], 0, price_max,
          list(np.arange(0, price_max + .01, .3)), '电价（元/kWh）')
    energy = [('负载', xs, stepped('load_kwh'), palette[0]), ('光伏', xs, stepped('pv_kwh'), palette[1]),
              ('购电', xs, stepped('purchase_kwh'), palette[2]), ('充电', xs, -stepped('charge_kwh'), palette[3]),
              ('放电', xs, stepped('discharge_kwh'), palette[4])]
    ymax = math.ceil(max(max(v) for _, _, v, _ in energy) / 500) * 500
    panel(can, 265, 225, energy, -1000, ymax, list(np.arange(-1000, ymax + 1, 500)), '时段电量（kWh）')
    panel(can, 25, 225, [('储电量', sx, soc, palette[0]), ('下限 1200', [0, 24], [1200, 1200], '#B34C46'),
                       ('上限 10800', [0, 24], [10800, 10800], '#8994A3')],
          0, 12000, [0, 3000, 6000, 9000, 12000], '储电量（kWh）', True)
    can.save()
    can = canvas.Canvas(str(FIG / 'problem1_soc.pdf'), pagesize=(730, 355), invariant=1)
    panel(can, 30, 300, [('储电量', sx, soc, palette[0]), ('下限 1200', [0, 24], [1200, 1200], '#B34C46'),
                       ('上限 10800', [0, 24], [10800, 10800], '#8994A3')],
          0, 12000, [0, 3000, 6000, 9000, 12000], '储电量（kWh）', True)
    can.save()
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    base, opt = float(summary['baseline_cost_yuan']), float(summary['dispatch_cost_yuan'])
    can = canvas.Canvas(str(FIG / 'problem1_cost_comparison.pdf'), pagesize=(610, 380), invariant=1)
    draw = Drawing(610, 380)
    bar = VerticalBarChart()
    bar.x, bar.y, bar.width, bar.height = 90, 65, 440, 260
    bar.data = [[base, opt]]
    bar.categoryAxis.categoryNames = ['无储能基线', '优化调度']
    bar.categoryAxis.labels.fontName = bar.valueAxis.labels.fontName = 'Chinese'
    bar.categoryAxis.labels.fontSize = bar.valueAxis.labels.fontSize = 11
    bar.valueAxis.valueMin, bar.valueAxis.valueMax = 0, math.ceil(base / 5000) * 5000 + 5000
    bar.valueAxis.valueStep = 10000
    bar.bars[0].fillColor = colors.HexColor(palette[0])
    bar.barLabels.fontName, bar.barLabels.fontSize = 'Chinese', 11
    bar.barLabelFormat = '%.2f'
    bar.barLabels.nudge = 10
    draw.add(bar)
    renderPDF.draw(draw, can, 0, 0)
    can.setFont('Chinese', 11)
    can.saveState()
    can.translate(28, 200)
    can.rotate(90)
    can.drawCentredString(0, 0, '全天购电费（元）')
    can.restoreState()
    can.drawCentredString(315, 20, f'节省 {base-opt:,.2f} 元（{(base-opt)/base:.2%}）')
    can.save()
    import pypdfium2 as pdfium
    validations = {}
    for pdf in sorted(FIG.glob('problem1_*.pdf')):
        reader = PdfReader(pdf)
        assert len(reader.pages) == 1
        page = reader.pages[0]
        assert '元' in page.extract_text() or '储电量' in page.extract_text()
        image_count = sum(1 for _ in page.images)
        assert image_count == 0, (pdf.name, 'not a pure vector PDF')
        doc = pdfium.PdfDocument(str(pdf))
        doc[0].render(scale=1.5).to_pil().save(QA / f'{pdf.stem}.png')
        doc.close()
        validations[pdf.name] = {'pages': 1, 'raster_images': image_count, 'size_bytes': pdf.stat().st_size}
    return validations


def table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |'] +
                     ['| ' + ' | '.join(str(v) for v in row) + ' |' for row in rows])


def make_report(audit, summaries, validation):
    a, b = summaries
    report = ['# 问题一计算结果报告', '', '本报告仅记录数据核验、模型计算与结果验收，不是论文正文。', '',
              '## 1. 输入核验', '',
              '已交叉核对项目中的《建模方案.md》、C题.pdf 的问题一、表1、表2、附录1与附件说明，以及官方结果模板。唯一数值输入为附件1.xlsx。', '',
              f"输入工作表：{audit['source_sheet_names']}；含表头维度：145 × 4；有效数据：144 × 3。表头为时间、电价、小区负载、光伏发电预测功率。",
              f"首时间 `{audit['first_time']}`，末时间 `{audit['last_time']}`；缺失 {len(audit['missing_cells'])}，非数值项 {len(audit['nonnumeric_cells'])}，重复时间 {audit['duplicate_time_count']}，重复整行 {audit['duplicate_full_row_count']}。所有数值有限且非负。", '',
              table(['变量', '最小值', '最大值'], [[k, v['min'], v['max']] for k, v in audit['numeric_ranges'].items()]), '',
              '光伏夜间零值和相同电价可能重复，这些是有效观测，不作为重复数据删除。原始文件的 SHA-256、软件版本与详细核验记录见 code/outputs/problem1_data_audit.json。', '',
              '## 2. 双口径时间映射与正式口径', '',
              'A：源时间是区间终点，00:10 → [00:00,00:10)，24:00 → [23:50,24:00)。',
              'B：源时间是区间起点，00:10 → [00:10,00:20)，24:00 → 次日 [00:00,00:10)。', '',
              f"模板共有144个标签，从 `{audit['template_first_label']}` 到 `{audit['template_last_label']}`。A在不跨日外推时仅精确覆盖143个；B精确覆盖144个。按用户选择规则，**正式口径采用B**。", '',
              '**B 的日界线解释：**题目一说明每天的电价和负载相同，本计算将这一天的光伏预测也作为待重复执行的典型日曲线。源文件最后一行的次日00:00—00:10，按24小时周期等价到日内00:00—00:10；其余各行按各自起始时刻进入00:00—24:00的优化时间轴。这个源行到民用时刻的显式连接记录在调度CSV中，并非按行号静默平移。', '',
              '因此，6000 kWh约束施加在真实日内00:00和24:00，而非00:10。正式解S(00:10)=6300 kWh。模板从00:10起展示，最后一行展示次日00:00—00:10，其购电量取周期重复的同一时段；模板顺序下的储电量从6300出发，经过24:00的6000，再回到次日00:10的6300。充放电汇总仍以题面要求的00:00—24:00民用日统计，每组恰含24段。没有把跨日段删去，也没有凭空增加第145个时段。', '',
              '这是一项明确的周期延拓假设，不能据此声称原题消除了时间歧义。若不接受周期延拓，B的00:10—次日00:10数据无法直接给出完整民用日的00:00—24:00汇总；需要补充日界观测或修改模板。A结果完整保留供对照。', '',
              table(['口径', '第一阶段最优费用/元', '第二阶段实际费用/元', '全天购电/kWh'],
                    [[s['mapping'], f"{s['optimal_cost_yuan']:.10f}", f"{s['dispatch_cost_yuan']:.10f}", f"{s['purchase_kwh']:.10f}"] for s in summaries]), '',
              f"B相对A的第一阶段最优费用差为 {b['optimal_cost_yuan']-a['optimal_cost_yuan']:.10f} 元；正式调度费用差为 {b['dispatch_cost_yuan']-a['dispatch_cost_yuan']:.10f} 元。成本很接近或相同也不意味着六个指定时段结果相同。", '',
              '## 3. 模型与求解状态', '',
              '所有功率先乘1/6 h转换为时段电量。模型为 min Σpₜqₜ；qₜ+vₜ+dₜ=ℓₜ+cₜ+wₜ；Sₜ₊₁=Sₜ+0.9cₜ−dₜ/0.9。储电量1200—10800 kWh，单段充放电上限5000/6 kWh，q与w非负，S(00:00)=S(24:00)=6000 kWh。变量c、d均按微网侧计量。', '',
              '不售电，富余电量可免费弃用；不添加购电容量上限、自放电、退化成本或网络损耗。题面90%在主模型解释为充、放电效率各90%，往返效率81%。', '',
              '每个口径721个连续变量、288条等式约束。第一阶段求最低购电费用；第二阶段在费用不超过J*+ε的前提下最小化总充放电量，ε=max(10⁻⁶,10⁻⁸|J*|)。费用限差是货币容差，不是能量平衡容差。程序含二进制互斥MILP回退；本次两个口径均未触发。', '',
              '若同一时段同时充放电，减少充电a及放电0.81a会保持储电变化不变，并增加可弃用余电0.19a，不增加购电费用，却降低总充放电量。因此在本题免费弃电且没有弃电上限的模型下，第二阶段可排除无效同时充放电。', '',
              table(['口径', '阶段1状态', '阶段2状态', 'ε/元', '费用上升/元', 'LP对偶间隙/元'],
                    [[s['mapping'], s['stage1_status'], s['stage2_status'], f"{s['epsilon_yuan']:.10g}",
                      f"{s['cost_increase_yuan']:.10g}", f"{s['lp_dual_gap_yuan']:.3g}"] for s in summaries]), '',
              f"正式求解方式：{b['solver']}。阶段1：{b['stage1_message']}；阶段2：{b['stage2_message']}。", '',
              '“最优费用”专指第一阶段全局LP最优值；正式模板保存第二阶段解，其真实费用用Σpₜqₜ重新计算，可能在指定ε内略高于第一阶段值。不可将两者混为同一精确数。', '',
              '## 4. 核心结果与无储能对照', '',
              table(['指标', '正式口径B'], [[name, f"{b[key]:,.10f}"] for name, key in [
                  ('正式调度购电费/元', 'dispatch_cost_yuan'), ('全天购电量/kWh', 'purchase_kwh'),
                  ('全天充电量/kWh', 'charge_kwh'), ('全天放电量/kWh', 'discharge_kwh'),
                  ('全天弃电量/kWh', 'curtailment_kwh'), ('无储能购电量/kWh', 'baseline_purchase_kwh'),
                  ('无储能购电费/元', 'baseline_cost_yuan'), ('节省额/元', 'saving_yuan')]]), '',
              f"无储能基线逐段购买 max(ℓ−v,0)，富余光伏弃用。正式调度费用节省比例为 **{b['saving_fraction']:.8%}**。", '',
              '### 表1：六个指定时段购电量', '']
    schedules = {m: read_csv(OUT / f'problem1_schedule_{m}.csv') for m in ['A', 'B']}
    lookups = {m: {r['time_label']: r for r in rows} for m, rows in schedules.items()}
    report += [table(['时间段', 'A购电量/kWh', 'B购电量/kWh（正式）'],
                     [[interval_label(h * 60)] + [f"{float(lookups[m][interval_label(h*60)]['purchase_kwh']):.10f}" for m in ['A', 'B']]
                      for h in [10, 12, 14, 16, 18, 20]]), '', '### 表2：六个四小时区间与首末储电量', '']
    blocks = read_csv(OUT / 'problem1_four_hour_summary.csv')
    report += [table(['口径', '时间段', '段数', '充电量/kWh', '放电量/kWh'],
                     [[r['mapping'], r['time_label'], r['interval_count'], f"{float(r['charge_kwh']):.10f}",
                       f"{float(r['discharge_kwh']):.10f}"] for r in blocks]), '',
               table(['口径', '00:00储电量/kWh', '24:00储电量/kWh'],
                     [[s['mapping'], s['soc_0000_kwh'], s['soc_2400_kwh']] for s in summaries]), '',
               '## 5. 校验结果', '',
               table(['指标', 'A', 'B（正式）'], [[k, f'{a[k]:.12g}' if isinstance(a[k], (int, float)) else a[k],
                                                 f'{b[k]:.12g}' if isinstance(b[k], (int, float)) else b[k]]
                     for k in ['energy_balance_max_abs_kwh', 'soc_transition_max_abs_kwh', 'soc_continuity_max_abs_kwh',
                               'soc_min_kwh', 'soc_max_kwh', 'charge_max_kwh', 'discharge_max_kwh',
                               'soc_terminal_abs_error_kwh', 'simultaneous_charge_discharge_max_kwh',
                               'block_charge_sum_abs_error_kwh', 'block_discharge_sum_abs_error_kwh',
                               'minimum_nonnegative_variable_kwh', 'nan_inf_count', 'period_count', 'blocks_each_24']]), '',
               '残差和上下界按10⁻⁶ kWh容差验收。物理上限使用精确5000/6=833.333333333…；题面833.333333是展示舍入值。', '',
               table(['输出验收项目', '结果'], [[k, v] for k, v in validation.items() if k != 'pdf_validation']), '',
               'Excel仅填144个购电结果、12个分组充放电结果和2个首末储电量，共158个原空白数值单元格。没有在模板空白处增加全天费用等摘要；这些值在本报告及summary.csv提供。原工作表名、行列结构、标签、格式及其他文件部件均保持不变。', '',
               '保存后独立重读CSV、Excel、PDF并检验逐项值与聚合。PDF均为单页矢量图，无嵌入位图；中文字体已嵌入。图数据均从已保存的正式schedule.csv与summary.csv读入。图中充电以负值绘制，只表示流向，CSV与Excel中的充电量仍非负。', '',
               '## 6. 文件与复现', '',
               '- `code/problem1.py`：从任意启动目录运行的入口，固定随机种子2026，主模型无随机性。',
               '- `code/fill_template.mjs`：模板数值写入与渲染辅助，由Python自动调用。为精确保留模板，导出的数值再回填到模板原始XML数值区域。',
               '- `result1.xlsx`：正式口径B结果。',
               '- `code/outputs/problem1_schedule.csv`：正式完整调度，按00:00—24:00排序，另列原模板标签与源Excel行号。',
               '- `code/outputs/problem1_summary.csv`：正式口径、费用、基线、节省与全部数值校验指标。',
               '- `code/outputs/problem1_schedule_A.csv`、`problem1_schedule_B.csv`、`problem1_mapping_comparison.csv`：双口径完整数值。',
               '- `code/outputs/problem1_time_mapping.csv`：每一源行、模板标签和日内时段之间的对应关系。',
               '- `code/outputs/problem1_six_periods.csv`、`problem1_four_hour_summary.csv`：表1、表2摘要。',
               '- `code/outputs/problem1_data_audit.json`、`problem1_validation.json`、`problem1_run.log`：输入核验、输出验收和运行日志。',
               '- `figures/problem1_dispatch.pdf`、`problem1_soc.pdf`、`problem1_cost_comparison.pdf`：三个中文矢量图。',
               '- `code/qa/`：模板及PDF渲染验收图。', '',
               '本机复现命令（PowerShell，启动目录不限）：', '', '```powershell',
               r'& "$env:USERPROFILE\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe" "D:\mycode\math\C题\code\problem1.py"',
               '```', '',
               '在具备依赖的其他环境：`python <项目根>/code/problem1.py`。Python依赖见code/requirements.txt；表格辅助需要Node与@oai/artifact-tool。默认自动寻找本机Codex捆绑运行时，可设置PROBLEM1_NODE；中文TrueType字体可设置PROBLEM1_CHINESE_FONT。程序不修改全局Python环境。', '',
               '交付时保留指定目录中的四份输入。所有计算路径均相对problem1.py确定，不依赖启动目录。问题二、三、四不在此次执行范围。', '']
    (REPORT / 'RESULTS_REPORT.md').write_text('\n'.join(report), encoding='utf-8')


def main():
    random.seed(2026)
    np.random.seed(2026)
    for directory in [OUT, QA, REPORT, FIG]:
        directory.mkdir(parents=True, exist_ok=True)
    numbers, times, labels, audit = audit_inputs()
    results = {m: solve(m, numbers, times, labels) for m in ['A', 'B']}
    summaries = [results[m][2] for m in ['A', 'B']]
    for m in ['A', 'B']:
        write_csv(OUT / f'problem1_schedule_{m}.csv', results[m][0])
    write_csv(OUT / 'problem1_schedule.csv', results['B'][0])
    all_blocks = results['A'][1] + results['B'][1]
    write_csv(OUT / 'problem1_four_hour_summary.csv', all_blocks)
    six = [{'mapping': m, 'time_label': r['time_label'], 'purchase_kwh': r['purchase_kwh']}
           for m in ['A', 'B'] for r in results[m][0] if int(r['civil_start_minute']) in [600, 720, 840, 960, 1080, 1200]]
    write_csv(OUT / 'problem1_six_periods.csv', six)
    by_source = {m: {r['source_excel_row']: r for r in results[m][0]} for m in ['A', 'B']}
    mapping = [{'source_excel_row': i + 2, 'source_time': clock_label(times[i]), 'template_label': labels[i],
                'A_interval': by_source['A'][i + 2]['time_label'],
                'B_civil_interval': by_source['B'][i + 2]['time_label'],
                'B_is_next_day_periodic_equivalent': times[i] == 1440} for i in range(144)]
    write_csv(OUT / 'problem1_time_mapping.csv', mapping)
    values = fill_template(results['B'][0], results['B'][1], labels)
    # Reopen saved data; verify physics again without using optimizer arrays.
    persisted = read_csv(OUT / 'problem1_schedule.csv')
    persisted_blocks = [r for r in read_csv(OUT / 'problem1_four_hour_summary.csv') if r['mapping'] == 'B']
    validate_schedule(persisted, persisted_blocks)
    validation = validate_excel(persisted, persisted_blocks, values, audit)
    # Recompute table 1 values and table 2 blocks from every saved schedule.
    for m in ['A', 'B']:
        data = read_csv(OUT / f'problem1_schedule_{m}.csv')
        lookup = {r['time_label']: r for r in data}
        for record in [r for r in read_csv(OUT / 'problem1_six_periods.csv') if r['mapping'] == m]:
            assert abs(float(record['purchase_kwh']) - float(lookup[record['time_label']]['purchase_kwh'])) <= 1e-8
        for group in [r for r in read_csv(OUT / 'problem1_four_hour_summary.csv') if r['mapping'] == m]:
            lo, hi = [minute_of(v) for v in group['time_label'].split('-')]
            selected = [r for r in data if lo <= int(r['civil_start_minute']) < hi]
            assert len(selected) == 24
            for k in ['charge_kwh', 'discharge_kwh']:
                assert abs(sum(float(r[k]) for r in selected) - float(group[k])) <= TOL
    validation['summary_tables_reaggregated'] = True
    persisted_cost = sum(float(r['price_yuan_per_kwh']) * float(r['purchase_kwh']) for r in persisted)
    persisted_base = sum(float(r['price_yuan_per_kwh']) * max(float(r['load_kwh']) - float(r['pv_kwh']), 0)
                         for r in persisted)
    validation['saved_cost_reaggregation_error_yuan'] = abs(persisted_cost - summaries[1]['dispatch_cost_yuan'])
    validation['saved_baseline_reaggregation_error_yuan'] = abs(persisted_base - summaries[1]['baseline_cost_yuan'])
    validation['saved_purchase_sum_error_kwh'] = abs(sum(float(r['purchase_kwh']) for r in persisted) - summaries[1]['purchase_kwh'])
    assert max(validation[k] for k in ['saved_cost_reaggregation_error_yuan', 'saved_baseline_reaggregation_error_yuan',
                                       'saved_purchase_sum_error_kwh']) <= TOL
    # Execute the template's cyclic order as an independent check of the midnight explanation.
    by_label = {r['template_time_label']: r for r in persisted}
    cyclic = [by_label[label] for label in labels]
    for before, after in zip(cyclic, cyclic[1:] + cyclic[:1]):
        assert abs(float(before['soc_end_kwh']) - float(after['soc_start_kwh'])) <= TOL
    validation['template_order_soc_cycle_closed'] = True
    summaries[1].update(validation)
    write_csv(OUT / 'problem1_summary.csv', [summaries[1]])
    write_csv(OUT / 'problem1_mapping_comparison.csv', [{k: s[k] for k in summaries[0]} for s in summaries])
    validation['pdf_validation'] = make_figures()
    node_run('final-preview')
    dump(OUT / 'problem1_validation.json', validation)
    make_report(audit, summaries, validation)
    for file in [ROOT / 'result1.xlsx', OUT / 'problem1_schedule.csv', OUT / 'problem1_summary.csv', REPORT / 'RESULTS_REPORT.md']:
        assert file.exists() and file.stat().st_size > 0
    text = (REPORT / 'RESULTS_REPORT.md').read_text(encoding='utf-8')
    assert all(x in text for x in ['10:00-10:10', '20:00-20:10', 'B', '校验'])
    print('ALL OUTPUT CHECKS PASSED\n' + json.dumps(summaries[1], ensure_ascii=False, indent=2), flush=True)


if __name__ == '__main__':
    class Tee:
        def __init__(self, stream, log):
            self.stream, self.log = stream, log
        def write(self, text):
            self.stream.write(text)
            self.log.write(text)
        def flush(self):
            self.stream.flush()
            self.log.flush()
    OUT.mkdir(parents=True, exist_ok=True)
    with (OUT / 'problem1_run.log').open('w', encoding='utf-8') as log:
        previous_out, previous_err = sys.stdout, sys.stderr
        sys.stdout, sys.stderr = Tee(previous_out, log), Tee(previous_err, log)
        try:
            main()
        except Exception:
            traceback.print_exc()
            raise SystemExit(1)
        finally:
            sys.stdout, sys.stderr = previous_out, previous_err
