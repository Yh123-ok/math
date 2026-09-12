"""Promote the validated simple causal MPC run into official files.

Run after simple_mpc.py: python 4test/experiments/finalize_mpc.py
"""
from collections import defaultdict
from pathlib import Path
import csv
import json
import os
import shutil
import subprocess
import sys

import numpy as np
from openpyxl import load_workbook

BASE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(BASE))
from src.data_io import read_inputs,sha256,interval_label
from src.reporting import emergency_events,find_node
from src.forecasting import generate_causal_forecasts
from src.simple_load_forecast import forecast_load

RUN=BASE/'experiments'/'simple_mpc_outputs'
OUT=BASE/'outputs'


def rows(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))


def main():
    summary=json.loads((RUN/'summary.json').read_text(encoding='utf-8'))
    data=read_inputs(BASE.parent)
    assert summary['input_sha256']=={k:v['sha256'] for k,v in data.audit.items()}
    sched=rows(RUN/'schedule.csv');daily=rows(RUN/'daily.csv')
    assert len(sched)==334*144 and len(daily)==334
    assert all(r['date']==str(data.dates[31+i]) for i,r in enumerate(daily))
    byday=defaultdict(list)
    for r in sched:byday[r['date']].append(r)
    plan_values=[];battery=[];events=[]
    for day,rr in byday.items():
        assert len(rr)==144 and [int(r['slot']) for r in rr]==list(range(144))
        q=np.array([float(r['plan_purchase_kwh']) for r in rr])
        c=np.array([float(r['charge_kwh']) for r in rr])
        u=np.array([float(r['discharge_kwh']) for r in rr])
        e=np.array([float(r['emergency_kwh']) for r in rr])
        p=np.array([float(r['price_yuan_per_kwh']) for r in rr])
        plan_values.append([*q.tolist(),float(q.sum()),float(q@p)])
        for b in range(6):
            sl=slice(b*24,(b+1)*24)
            battery.append({'date':day,'time_range':f'{4*b:02d}:00-{4*b+4:02d}:00',
                'charge_kwh':float(c[sl].sum()),'discharge_kwh':float(u[sl].sum()),
                'time':'00:00' if b==0 else ('24:00' if b==1 else ''),
                'soc_kwh':float(rr[0]['soc_start_kwh']) if b==0 else
                   (float(rr[-1]['soc_end_kwh']) if b==1 else '')})
        events.extend({'date':day,'time_range':period,'emergency_kwh':amount}
            for period,amount in emergency_events(e))
    old=json.loads((BASE/'archives'/'joint_v1'/'summary.json').read_text(encoding='utf-8'))
    assert summary['total_cost_yuan'] < old['total_cost_yuan']
    payload={'output_path':str(BASE/'result4-2.xlsx'),'plan_values':plan_values,
             'battery_rows':battery,'emergency_rows':events,
             'template_labels':list(data.template_labels),
             'physical_labels':[interval_label(t) for t in range(144)]}
    OUT.mkdir(exist_ok=True)
    payload_path=OUT/'mpc_excel_payload.json'
    payload_path.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
    env=os.environ.copy()
    env['PROBLEM4_PAYLOAD']=str(payload_path)
    env['PROBLEM4_OUTPUT']=str(BASE/'result4-2.xlsx')
    env['PROBLEM4_EXPORT_RECEIPT']=str(OUT/'mpc_workbook_export.json')
    env['PROBLEM4_QA_DIR']=str(OUT/'qa_mpc')
    subprocess.run([find_node(BASE),str(BASE/'tools'/'write_result.mjs')],cwd=BASE,env=env,check=True)
    receipt=json.loads((OUT/'mpc_workbook_export.json').read_text(encoding='utf-8'))
    assert Path(receipt['output_path']).resolve()==(BASE/'result4-2.xlsx').resolve(),receipt
    original=load_workbook(BASE.parent/'附件'/'附件5'/'result4-2.xlsx',data_only=True)
    result=load_workbook(BASE/'result4-2.xlsx',data_only=True)
    assert original.sheetnames==result.sheetnames
    for a,b in zip(original.worksheets,result.worksheets):
        assert [a.cell(1,j).value for j in range(1,a.max_column+1)]==[
               b.cell(1,j).value for j in range(1,a.max_column+1)],a.title
    errors=[]
    for i,expected in enumerate(plan_values,2):
        assert result.worksheets[0].cell(i,1).value.date()==data.dates[31+i-2]
        for j,v in enumerate(expected,2):
            errors.append(abs(float(result.worksheets[0].cell(i,j).value)-v))
    assert max(errors)<1e-6
    battery_error=0.
    for i,r in enumerate(battery,2):
        ws=result['充放电量']
        assert ws.cell(i,2).value==r['time_range']
        assert ws.cell(i,1).value.date().isoformat()==r['date']
        battery_error=max(battery_error,abs(float(ws.cell(i,3).value)-r['charge_kwh']),
            abs(float(ws.cell(i,4).value)-r['discharge_kwh']))
        assert ws.cell(i,5).value in (r['time'],None if r['time']=='' else r['time'])
        if r['soc_kwh']!='':battery_error=max(battery_error,
            abs(float(ws.cell(i,6).value)-r['soc_kwh']))
    assert battery_error<1e-6
    emergency_error=0.
    for i,r in enumerate(events,2):
        ws=result['紧急购电量']
        assert ws.cell(i,2).value==r['time_range']
        assert ws.cell(i,1).value.date().isoformat()==r['date']
        emergency_error=max(emergency_error,abs(float(ws.cell(i,3).value)-r['emergency_kwh']))
    assert emergency_error<1e-6
    original.close();result.close()
    assert sha256(BASE.parent/'附件'/'附件5'/'result4-2.xlsx')==data.audit['template']['sha256']
    perday_error=0.
    for r in daily:
        rr=byday[r['date']]
        def total(col):return sum(float(x[col]) for x in rr)
        normal=sum(float(x['price_yuan_per_kwh'])*float(x['plan_purchase_kwh']) for x in rr)
        urgent=sum(5*float(x['price_yuan_per_kwh'])*float(x['emergency_kwh']) for x in rr)
        checks={'plan_purchase_kwh':total('plan_purchase_kwh'),'emergency_kwh':total('emergency_kwh'),
            'charge_kwh':total('charge_kwh'),'discharge_kwh':total('discharge_kwh'),
            'waste_kwh':total('waste_kwh'),'planned_cost_yuan':normal,
            'emergency_cost_yuan':urgent,'total_cost_yuan':normal+urgent,
            'soc_start_kwh':float(rr[0]['soc_start_kwh']),
            'soc_end_kwh':float(rr[-1]['soc_end_kwh'])}
        perday_error=max(perday_error,max(abs(float(r[k])-v) for k,v in checks.items()))
    assert perday_error<1e-6
    assert abs(sum(float(r['total_cost_yuan']) for r in daily)-summary['total_cost_yuan'])<1e-6
    assert all(np.isfinite(float(r[k])) for r in sched for k in
        ('price_yuan_per_kwh','load_kwh','pv_kwh','plan_purchase_kwh','charge_kwh',
         'discharge_kwh','emergency_kwh','waste_kwh','soc_start_kwh','soc_end_kwh'))
    assert all(float(r[k])>=-1e-6 for r in sched for k in
        ('plan_purchase_kwh','charge_kwh','discharge_kwh','emergency_kwh','waste_kwh'))
    for src,name in ((RUN/'schedule.csv','official_schedule.csv'),
                     (RUN/'daily.csv','official_daily.csv'),
                     (RUN/'summary.json','official_summary.json')):
        shutil.copy2(src,OUT/name)
    original_forecasts=generate_causal_forecasts(data.dates,data.load,data.pv)
    improved_load=forecast_load(data.dates,data.load)
    with (OUT/'official_forecasts.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['date','slot','period','actual_load_kwh','old_load_forecast_kwh',
            'new_load_forecast_kwh','actual_pv_kwh','pv_forecast_kwh'])
        for d in range(31,len(data.dates)):
            for t in range(144):w.writerow([str(data.dates[d]),t,interval_label(t),data.load[d,t],
                original_forecasts['load_final'][d,t],improved_load[d,t],data.pv[d,t],
                original_forecasts['pv_final'][d,t]])
    stale=OUT/'official_selections.csv'
    if stale.exists():stale.unlink()
    audit={'original_template_headers_unchanged':True,'plan_values_checked':len(errors),
           'max_excel_plan_error_kwh':max(errors),'battery_rows':len(battery),
           'max_excel_battery_error_kwh':battery_error,'emergency_event_rows':len(events),
           'max_excel_emergency_error_kwh':emergency_error,'template_sha256_unchanged':True,
           'max_csv_daily_reaggregation_error':perday_error,
           'all_physics':summary['checks']}
    (OUT/'official_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in audit.items() if k!='all_physics'},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
