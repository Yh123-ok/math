"""保存调度、汇总和模板载荷；从落盘数据绘图并独立重读Excel验收。"""
import csv, json, os, shutil, subprocess, sys
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
import config as cfg
from .data_io import interval_label, sha256


def write_csv(path, rows):
    rows=list(rows)
    if not rows:return
    with path.open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)


def read_csv(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))


def emergency_events(values):
    events=[];start=None
    for t in range(145):
        positive=t<144 and values[t]>1e-8
        if positive and start is None:start=t
        if not positive and start is not None:
            a,b=start*10,t*10
            events.append((f'{a//60:02d}:{a%60:02d}-{b//60:02d}:{b%60:02d}',float(values[start:t].sum())))
            start=None
    return events


def find_node(base):
    """支持普通Node环境，以及Codex提供的离线运行库；不改系统依赖。"""
    runtime=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    node=os.environ.get('PROBLEM4_NODE') or shutil.which('node')
    if not node and (runtime/'bin/node.exe').exists():node=str(runtime/'bin/node.exe')
    if not node:raise RuntimeError('需安装Node.js并提供@oai/artifact-tool，或设置PROBLEM4_NODE')
    target=base/'node_modules'
    if not target.exists() and (runtime/'node_modules').exists():
        if os.name=='nt':
            # 原生PowerShell创建目录联接，不安装或修改依赖包。
            def quote(p):return "'"+str(p).replace("'","''")+"'"
            subprocess.run(['powershell','-NoProfile','-Command',
                f'New-Item -ItemType Junction -Path {quote(target)} -Target {quote(runtime/"node_modules")} | Out-Null'],check=True)
        else:target.symlink_to(runtime/'node_modules',target_is_directory=True)
    return node


def export_results(base,data,forecasts,streams,daily,selections,risk_history,checks,make_artifacts=True):
    out=base/'outputs';(base/'reports').mkdir(exist_ok=True)
    dispatch=[];forecast_rows=[];battery=[];events=[];plan_values=[]
    for d,record in enumerate(streams['official']):
        day=str(data.dates[d]);q=record['plan']['q'];a=record['actual']
        for t in range(144):
            dispatch.append({'date':day,'slot':t,'period':interval_label(t),
                'price_yuan_per_kwh':data.prices[d,t],'load_kwh':data.load[d,t],'pv_kwh':data.pv[d,t],
                'plan_purchase_kwh':q[t],**{k+'_kwh':a[k][t] for k in ('charge','discharge','emergency','waste')},
                'soc_start_kwh':a['soc'][t],'soc_end_kwh':a['soc'][t+1]})
            if d:
                forecast_rows.append({'date':day,'slot':t,'period':interval_label(t),
                    'load_forecast_kwh':forecasts['load_final'][d,t],'pv_forecast_kwh':forecasts['pv_final'][d,t],
                    'net_point_forecast_kwh':record['point_forecast'][t],
                    'net_plan_demand_kwh':record['demand'][t],
                    'net_error_kwh':forecasts['net_error'][d,t]})
        if d<cfg.OFFICIAL_START_INDEX:continue
        plan_values.append([*q.tolist(),float(q.sum()),daily['official'][d]['planned_cost_yuan']])
        for b in range(6):
            sl=slice(b*24,(b+1)*24)
            battery.append({'date':day,'time_range':f'{4*b:02d}:00-{4*b+4:02d}:00',
                'charge_kwh':float(a['charge'][sl].sum()),'discharge_kwh':float(a['discharge'][sl].sum()),
                'time':'00:00' if b==0 else ('24:00' if b==1 else ''),
                'soc_kwh':float(a['soc'][0]) if b==0 else (float(a['soc'][-1]) if b==1 else '')})
        events.extend({'date':day,'time_range':label,'emergency_kwh':value} for label,value in emergency_events(a['emergency']))
    write_csv(out/'dispatch_all_year.csv',dispatch)
    write_csv(out/'forecasts.csv',forecast_rows)
    write_csv(out/'battery_four_hour.csv',battery)
    write_csv(out/'emergency_events.csv',events)
    write_csv(out/'selections.csv',selections)
    write_csv(out/'time_mapping.csv',({'slot':t,'source_endpoint':data.source_times[t],
        'original_template_label':data.template_labels[t],'output_template_label':data.template_labels[t],
        'physical_interval':interval_label(t),'source_data_column_1based':t+2,
        'output_excel_column_1based':t+2} for t in range(144)))
    forecast_weights=[]
    for d in range(1,len(data.dates)):
        forecast_weights.append({'date':str(data.dates[d]),'load_weight_a':forecasts['load_weight_a'][d],
            'net_ridge_weight':forecasts['net_ridge_weight'][d],
            **{f'pv_weight_{k}':forecasts['pv_weights'][d,j] for j,k in enumerate('abc')}})
    write_csv(out/'forecast_weights.csv',forecast_weights)
    daily_rows=[];summary=[]
    sum_fields=('plan_purchase_kwh','emergency_kwh','charge_kwh','discharge_kwh','waste_kwh',
                'planned_cost_yuan','emergency_cost_yuan','total_cost_yuan')
    for name,rows in daily.items():
        daily_rows.extend({'strategy':name,'date':str(data.dates[r['day']]),**r} for r in rows)
        selected=rows[cfg.OFFICIAL_START_INDEX:]
        summary.append({'strategy':name,**{k:sum(r[k] for r in selected) for k in sum_fields},
            'soc_start_kwh':selected[0]['soc_start_kwh'],'soc_end_kwh':selected[-1]['soc_end_kwh']})
    write_csv(out/'daily_summary.csv',daily_rows)
    write_csv(out/'summary.csv',summary)
    write_csv(out/'risk_shadow_daily.csv',({'risk_index':j,'date':str(data.dates[r['day']]),**r}
        for j,rows in enumerate(risk_history) for r in rows))
    payload={'output_path':str(base/'result4-2.xlsx'),'plan_values':plan_values,'battery_rows':battery,
        'emergency_rows':events,'template_labels':list(data.template_labels),
        'physical_labels':[interval_label(t) for t in range(144)]}
    (out/'excel_payload.json').write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
    if not make_artifacts:
        write_report(base,data,summary,daily,streams,checks)
        print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
        return
    node=find_node(base)
    subprocess.run([node,str(base/'tools/write_result.mjs')],check=True,cwd=base)
    exported=Path(json.loads((out/'workbook_export.json').read_text(encoding='utf-8'))['output_path'])
    checks['saved_outputs']=verify_saved_outputs(base,data,payload,summary,exported)
    from .figures import create_figures
    create_figures(base)
    (out/'checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    write_report(base,data,summary,daily,streams,checks)
    import scipy,reportlab,openpyxl
    artifact_paths=[p for folder in ('src','tools','reports','figures','outputs')
        for p in (base/folder).rglob('*') if p.is_file()]
    artifact_paths += [p for p in base.iterdir() if p.is_file()]
    manifest={'python':sys.version,'numpy':np.__version__,'scipy':scipy.__version__,
        'reportlab':reportlab.Version,'openpyxl':openpyxl.__version__,
        'config':{k:v for k,v in vars(cfg).items() if k.isupper()},'inputs':data.audit,
        'outputs_sha256':{str(p.relative_to(base)):sha256(p) for p in artifact_paths
            if p.suffix in ('.csv','.pdf','.xlsx','.md','.py','.mjs')}}
    (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)


def verify_saved_outputs(base,data,payload,summary,workbook_path=None):
    out=base/'outputs';rows=read_csv(out/'dispatch_all_year.csv')
    assert len(rows)==365*144
    numeric=list(rows[0])[3:]
    x=np.array([[float(r[k]) for k in numeric] for r in rows])
    assert np.isfinite(x).all()
    official=rows[31*144:];q=np.array([float(r['plan_purchase_kwh']) for r in official]).reshape(334,144)
    prices=np.array([float(r['price_yuan_per_kwh']) for r in official]).reshape(334,144)
    urgent=np.array([float(r['emergency_kwh']) for r in official]).reshape(334,144)
    workbook_path=workbook_path or base/'result4-2.xlsx'
    wb=load_workbook(workbook_path,data_only=True)
    assert wb.sheetnames==['计划购电量','充放电量','紧急购电量']
    ws=wb.worksheets[0]
    assert [ws.cell(1,t+2).value for t in range(144)]==payload['template_labels']
    original=load_workbook(base.parent/'附件/附件5/result4-2.xlsx',read_only=True,data_only=True)
    header_cells=0
    for source,filled in zip(original,wb):
        for col in range(1,source.max_column+1):
            assert source.cell(1,col).value==filled.cell(1,col).value,(source.title,col,'表头与原模板不一致')
            header_cells+=1
    original.close()
    mapping=read_csv(out/'time_mapping.csv')
    assert len(mapping)==144
    for t,r in enumerate(mapping):
        assert r['output_template_label']==data.template_labels[t]
        assert r['source_endpoint']==data.source_times[t] and r['physical_interval']==interval_label(t)
        assert int(r['source_data_column_1based'])==int(r['output_excel_column_1based'])==t+2
    excel=np.array([[ws.cell(d+2,t+2).value for t in range(146)] for d in range(334)],float)
    expected=np.c_[q,q.sum(axis=1),(q*prices).sum(axis=1)]
    error=float(np.max(abs(excel-expected)));assert error<1e-6
    for d in range(334):assert ws.cell(d+2,1).value.date()==data.dates[d+31]
    b=wb.worksheets[1];berror=0.
    battery_csv_error=0.
    for j,r in enumerate(read_csv(out/'battery_four_hour.csv')):
        d,block=divmod(j,6)
        part=official[d*144+block*24:d*144+(block+1)*24]
        assert len(part)==24
        for key in ('charge_kwh','discharge_kwh'):
            battery_csv_error=max(battery_csv_error,abs(sum(float(x[key]) for x in part)-float(r[key])))
    assert battery_csv_error<1e-6
    for i,r in enumerate(payload['battery_rows'],2):
        assert b.cell(i,1).value.date().isoformat()==r['date'] and b.cell(i,2).value==r['time_range']
        for col,key in ((3,'charge_kwh'),(4,'discharge_kwh'),(6,'soc_kwh')):
            if r[key]!='':berror=max(berror,abs(float(b.cell(i,col).value)-r[key]))
        assert (b.cell(i,5).value or '')==r['time']
    e=wb.worksheets[2];eerror=0.
    for i,r in enumerate(payload['emergency_rows'],2):
        assert e.cell(i,1).value.date().isoformat()==r['date'] and e.cell(i,2).value==r['time_range']
        eerror=max(eerror,abs(float(e.cell(i,3).value)-r['emergency_kwh']))
    assert max(berror,eerror)<1e-6
    wb.close()
    event_sum=sum(float(r['emergency_kwh']) for r in read_csv(out/'emergency_events.csv'))
    assert abs(event_sum-urgent.sum())<1e-6
    # 重新聚合原始十分钟CSV的全部费用，不读取内存日汇总。
    official_summary=next(r for r in summary if r['strategy']=='official')
    cost_error=abs(float((prices*(q+5*urgent)).sum())-official_summary['total_cost_yuan'])
    assert cost_error<1e-6
    assert sha256(data_path:=base.parent/'附件/附件5/result4-2.xlsx')==data.audit['template']['sha256']
    return {'verified_workbook':str(workbook_path.relative_to(base)),
        'original_header_cells_verified':header_cells,'original_headers_unchanged':True,
        'time_mapping_rows_verified':144,'excel_plan_and_daily_totals_max_error':error,'excel_battery_max_error':berror,
        'csv_battery_reaggregation_error_kwh':battery_csv_error,
        'excel_emergency_max_error':eerror,'csv_cost_reaggregation_error_yuan':cost_error,
        'emergency_event_sum_error_kwh':abs(event_sum-urgent.sum()),'template_unchanged':True,
        'official_days':334,'official_purchase_values':334*144,'battery_rows':len(payload['battery_rows']),
        'emergency_event_rows':len(payload['emergency_rows'])}


def write_report(base,data,summary,daily,streams,checks):
    text=['# 第四问重解第二问：计算结果','',
        '评价期为2025-02-01至2025-12-31；1月也逐日运行并延续真实模拟库存。以下为历史因果回测，不是全信息最优解，也不是独立未触碰测试集。',
        '', '假设每天0:00已公布当天144段电价；未知明日电价。附件2源时间作为区间终点。按用户要求，Excel所有表头原样保留。计划表第t个数据列对应源数据第t个终点及其前十分钟物理区间，不按模板偏移的文字平移数值；详见time_mapping.csv。附件3不进入本问题的信息集。',
        '', '## 1. 核心结果与对照','',
        '|策略|正常购电量/kWh|紧急购电量/kWh|正常费用/元|五倍紧急费用/元|总费用/元|',
        '|---|---:|---:|---:|---:|---:|']
    labels={'official':'正式：历史选择控制器','greedy':'对照：即时放电','mpc':'对照：价格感知MPC'}
    for r in summary:
        text.append('|'+labels[r['strategy']]+'|'+'|'.join(f'{r[k]:.6f}' for k in ('plan_purchase_kwh','emergency_kwh','planned_cost_yuan','emergency_cost_yuan','total_cost_yuan'))+'|')
    text.extend(['','正式方案始终按历史分数决定，不根据年底总费用倒选全年赢家。三条策略均采用同一历史风险选择序列，但库存分别连续演化，故购电计划可不同。','',
        '## 2. 题面四个指定日期','',
        '|日期|10:00|12:00|14:00|16:00|18:00|20:00|全天计划量/kWh|计划费/元|紧急量/kWh|紧急费/元|',
        '|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|'])
    for date in cfg.SPECIFIED_DATES:
        d=next(i for i,x in enumerate(data.dates) if str(x)==date);r=daily['official'][d];q=streams['official'][d]['plan']['q']
        values=[q[t] for t in (60,72,84,96,108,120)]+[r[k] for k in ('plan_purchase_kwh','planned_cost_yuan','emergency_kwh','emergency_cost_yuan')]
        text.append('|'+date+'|'+'|'.join(f'{v:.6f}' for v in values)+'|')
    text.extend(['','上表时刻表示其后十分钟区间，例如10:00表示10:00–10:10。','',
        '|日期|四小时区间|充电/kWh|放电/kWh|','|---|---|---:|---:|'])
    for r in read_csv(base/'outputs/battery_four_hour.csv'):
        if r['date'] in cfg.SPECIFIED_DATES:text.append(f"|{r['date']}|{r['time_range']}|{float(r['charge_kwh']):.6f}|{float(r['discharge_kwh']):.6f}|")
    text.extend(['','|日期|0:00储电量/kWh|24:00储电量/kWh|','|---|---:|---:|'])
    for r in daily['official']:
        date=str(data.dates[r['day']])
        if date in cfg.SPECIFIED_DATES:text.append(f"|{date}|{r['soc_start_kwh']:.6f}|{r['soc_end_kwh']:.6f}|")
    text.extend(['','|日期|紧急购电区间|紧急量/kWh|','|---|---|---:|'])
    events=read_csv(base/'outputs/emergency_events.csv')
    for date in cfg.SPECIFIED_DATES:
        found=[r for r in events if r['date']==date]
        if not found:text.append(f'|{date}|无|0.000000|')
        for r in found:text.append(f"|{date}|{r['time_range']}|{float(r['emergency_kwh']):.6f}|")
    text.extend(['','四日及全年所有紧急时段见 `outputs/emergency_events.csv` 与Excel第三页；四小时区间均恰含24段。','',
        '## 3. 数据、求解与验收','',
        '附件2两页及附件4一页均366行×145列（含表头），365日×144段。逐项检查日期、时间、缺失、非数值、有限性及范围，详情如下。',
        '', '```json',json.dumps(data.audit,ensure_ascii=False,indent=2),'```','',
        '所有日前两阶段LP和日内LP均由SciPy HiGHS成功求解；任何失败或残差越界都会终止程序。验收指标如下。',
        '', '```json',json.dumps(checks,ensure_ascii=False,indent=2),'```','',
        '扰动测试覆盖2月、6月、12月的当前/未来实际值；另检查未来电价、日内未来实际值、未来调参分数。测试通过支持这些已覆盖路径的因果性，不代表对所有可能程序路径的形式证明。',
        '', '## 4. 方法和复现','',
        '完整流程、公式和选择依据见 `METHOD.md`，运行命令及每个代码文件用途见上级 `README.md`。PDF均由已保存CSV生成。',
        '总费用只含实际结算的计划购电费用与五倍紧急购电费用；软终端库存惩罚仅用于决策，不计入账单。'])
    (base/'reports/RESULTS_REPORT.md').write_text('\n'.join(text)+'\n',encoding='utf-8')
