"""只读附件2、4和结果模板，按区间终点建立唯一物理时间轴。"""
from dataclasses import dataclass
from datetime import date, timedelta, time
from pathlib import Path
import hashlib
import numpy as np
from openpyxl import load_workbook


def interval_label(t):
    a, b = 10*t, 10*(t+1)
    return f'{a//60:02d}:{a%60:02d}-{b//60:02d}:{b%60:02d}'


def clock_text(value):
    return value.strftime('%H:%M') if isinstance(value, time) else str(value).strip()


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


@dataclass
class Data:
    dates: tuple
    load: np.ndarray
    pv: np.ndarray
    prices: np.ndarray
    source_times: tuple
    template_labels: tuple
    audit: dict


def read_inputs(project):
    dates = tuple(date(2025,1,1)+timedelta(days=i) for i in range(365))
    times = tuple([f'{m//60:02d}:{m%60:02d}' for m in range(10,1440,10)]+['0:00+1'])
    audit, arrays = {}, []
    for file, sheets in [('附件2.xlsx', ['小区负载','光伏发电实际功率']),
                         ('附件4.xlsx', ['Sheet1'])]:
        path = project/'附件'/file
        wb = load_workbook(path, read_only=True, data_only=True)
        assert wb.sheetnames == sheets, (file, wb.sheetnames)
        info = {'sha256':sha256(path), 'sheets':{}}
        for name in sheets:
            ws = wb[name]
            rows = list(ws.values)
            assert (len(rows),len(rows[0])) == (366,145)
            assert tuple(clock_text(x) for x in rows[0][1:]) == times
            found = tuple(r[0].date() for r in rows[1:])
            assert found == dates
            x = np.array([r[1:] for r in rows[1:]],dtype=float)
            assert np.isfinite(x).all() and (x>=0).all()
            if file=='附件4.xlsx': assert (x>0).all(), '当前策略假设电价为正'
            arrays.append(x if file=='附件4.xlsx' else x/6)
            info['sheets'][name] = {'shape':[366,145], 'header_first':str(rows[0][0]),
                'first_date':str(found[0]),'last_date':str(found[-1]),
                'first_time':times[0],'last_time':times[-1], 'missing_non_numeric_nonfinite':0,
                'duplicate_dates':len(found)-len(set(found)), 'min':float(x.min()),'max':float(x.max())}
        audit[file]=info
        wb.close()
    template = project/'附件/附件5/result4-2.xlsx'
    wb = load_workbook(template,read_only=False,data_only=True)
    assert wb.sheetnames==['计划购电量','充放电量','紧急购电量']
    labels=tuple(wb.worksheets[0].cell(1,c).value for c in range(2,146))
    assert tuple(wb.worksheets[0].cell(r,1).value.date() for r in range(2,336))==dates[31:]
    audit['template']={'sha256':sha256(template),'sheets':{w.title:[w.max_row,w.max_column] for w in wb},
        'label_correction':'output copy only: source endpoint 00:10 maps to 00:00-00:10; no array rotation'}
    wb.close()
    return Data(dates,*arrays,times,labels,audit)
